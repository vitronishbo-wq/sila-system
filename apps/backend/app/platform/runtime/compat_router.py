from __future__ import annotations

import asyncio
import json
import os
import time
import uuid
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any

import yaml
from apps.backend.app.api.deps import get_current_user, get_db
from apps.backend.app.core.audit import SLA_DEFINITIONS, evaluate_sla_status, get_sla_for_service
from apps.backend.app.core.db import AsyncSessionLocal
from apps.backend.app.core.observability import Metrics
from apps.backend.app.core.observability.otel_integration import OTEL_AVAILABLE
from apps.backend.app.core.settings import settings
from apps.backend.app.modules.economy.application.dto.invoice_schema import CreateInvoiceSchema
from apps.backend.app.modules.economy.application.dto.payment_schema import CreatePaymentSchema
from apps.backend.app.modules.economy.core.application.services.invoice_service import InvoiceService
from apps.backend.app.modules.economy.core.application.services.payment_service import PaymentService
from apps.backend.app.modules.economy.domain.exceptions import (
    DomainValidationError,
    DuplicatePaymentError,
    FUCError,
    InvalidInvoiceStateError,
    InvoiceNotFoundError,
)
from apps.backend.app.modules.economy.domain.models.invoice import Invoice
from apps.backend.app.modules.economy.domain.models.payment import Payment
from apps.backend.app.modules.economy.infrastructure.adapters import (
    SQLAlchemyInvoiceRepository,
    SQLAlchemyPaymentRepository,
)
from core.security import verify_password
from fastapi import (
    APIRouter,
    Depends,
    Header,
    HTTPException,
    Query,
    Request,
    Response,
    WebSocket,
    WebSocketDisconnect,
    status,
)
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import String, bindparam, text
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.responses import StreamingResponse

from apps.backend.core.auth import JWTHandler

router = APIRouter()
_MODULES_PATH = Path(__file__).resolve().parents[2] / "modules"
_MODULE_LABELS = {
    "identity": "Identificacao Civil",
    "justice": "Justica e Registo",
    "registry": "Registo Civil",
    "economy": "Financas",
    "tax": "Financas",
    "energy": "Energia",
    "water": "Agua e Saneamento",
    "employment": "Emprego e Trabalho",
    "society": "Emprego e Assistencia Social",
    "logistics": "Transportes",
    "licensing": "Licenciamento",
    "transport": "Transportes",
    "notaries": "Cartorios e Notariado",
    "saude": "Saude",
    "educacao": "Educacao",
}
_CATEGORY_MODULE_MAP = {
    "tax": "economy",
    "registry": "justice",
    "notaries": "justice",
    "transport": "logistics",
    "employment": "society",
}
_MODULE_UI_OVERRIDES = {
    "identity": {"icon": "fa-id-card", "color": "#2563eb", "order": 1},
    "registry": {"icon": "fa-book", "color": "#10b981", "order": 2},
    "justice": {"icon": "fa-scale-balanced", "color": "#0ea5e9", "order": 3},
    "tax": {"icon": "fa-receipt", "color": "#f59e0b", "order": 4},
    "economy": {"icon": "fa-landmark", "color": "#f59e0b", "order": 4},
    "water": {"icon": "fa-tint", "color": "#06b6d4", "order": 5},
    "energy": {"icon": "fa-bolt", "color": "#f97316", "order": 6},
    "employment": {"icon": "fa-briefcase", "color": "#8b5cf6", "order": 7},
    "society": {"icon": "fa-briefcase", "color": "#8b5cf6", "order": 7},
    "educacao": {"icon": "fa-graduation-cap", "color": "#22c55e", "order": 6},
    "saude": {"icon": "fa-heart-pulse", "color": "#ef4444", "order": 5},
    "licensing": {"icon": "fa-file-signature", "color": "#ef4444", "order": 8},
    "transport": {"icon": "fa-bus", "color": "#22c55e", "order": 9},
    "logistics": {"icon": "fa-bus", "color": "#22c55e", "order": 9},
    "notaries": {"icon": "fa-stamp", "color": "#6366f1", "order": 10},
}
_SERVICES: list[dict[str, Any]] = [
    {"id": "identity", "code": "BI_EMISSAO", "name": "Bilhete de Identidade", "price": 1000.0},
    {
        "id": "registry",
        "code": "CERTIDAO_REGISTO",
        "name": "Certidão do Registo Civil",
        "price": 500.0,
    },
    {"id": "tax", "code": "NIF_EMISSAO", "name": "Emissão de NIF", "price": 750.0},
    {"id": "water", "code": "AGUA_LIGACAO", "name": "Ligação de Água", "price": 3500.0},
    {"id": "energy", "code": "ENERGIA_LIGACAO", "name": "Ligação de Energia", "price": 3500.0},
    {
        "id": "employment",
        "code": "EMPREGO_CANDIDATURA",
        "name": "Candidatura a Emprego",
        "price": 0.0,
    },
    {"id": "licensing", "code": "LICENCIAMENTO", "name": "Pedido de Licença", "price": 2500.0},
    {
        "id": "transport",
        "code": "TRANSPORTES_REGISTO",
        "name": "Licenciamento de Transportes",
        "price": 3000.0,
    },
    {"id": "notaries", "code": "CERTIDAO_NOTARIAL", "name": "Certidão Notarial", "price": 1200.0},
]
_SERVICE_CATALOG: list[dict[str, Any]] = [
    {
        "code": "BI_EMISSAO",
        "name": "Bilhete de Identidade",
        "category": "identity",
        "description": "Emissao e renovacao do bilhete de identidade",
        "price": 1000.0,
    },
    {
        "code": "CERTIDAO_REGISTO",
        "name": "Certidao do Registo Civil",
        "category": "registry",
        "description": "Emissao de certidoes de registo civil",
        "price": 500.0,
    },
    {
        "code": "NIF_EMISSAO",
        "name": "Emissao de NIF",
        "category": "tax",
        "description": "Cadastro e regularizacao fiscal",
        "price": 750.0,
    },
]
_ORDERS: dict[str, dict[str, Any]] = {}
_PAYMENTS: dict[str, dict[str, Any]] = {}
_PING_MODULES = {
    "urbanism",
    "justice",
    "commercial",
    "education",
    "address",
    "registry",
    "reports",
    "identity",
    "social",
    "common",
    "governance",
    "journeys",
    "services",
    "training",
    "internal",
}
_EXPORT_TTL_SECONDS = 3600


class OrderCreatePayload(BaseModel):
    service_id: str = Field(..., min_length=1)


class DocumentPayload(BaseModel):
    filename: str
    content_type: str
    size_bytes: int
    uri: str


class DocumentsPayload(BaseModel):
    documents: list[DocumentPayload]


class BIEmitPayload(BaseModel):
    citizen_fuc_id: str


class CatalogField(BaseModel):
    key: str
    label: str
    type: str = "text"
    required: bool = False
    options: list[str] | None = None
    placeholder: str | None = None
    mask: str | None = None
    helper: str | None = None
    multiple: bool | None = None


class CatalogForm(BaseModel):
    model_config = ConfigDict(populate_by_name=True)
    code: str
    form_schema: dict[str, Any] = Field(alias="schema")


class InvoiceResponse(BaseModel):
    id: str
    citizen_id: str
    reference: str
    revenue_code: str
    cost_center: str
    service_code: str
    service_name: str
    amount: float
    currency: str
    status: str
    created_at: str
    updated_at: str
    due_date: str


class PaymentResponse(BaseModel):
    id: str
    invoice_id: str
    citizen_id: str
    amount: float
    currency: str
    gateway_reference: str
    payment_method: str
    status: str
    created_at: str
    confirmation_timestamp: str | None = None


def _get_financas_services(db: AsyncSession):
    invoice_repo = SQLAlchemyInvoiceRepository(db)
    payment_repo = SQLAlchemyPaymentRepository(db)
    invoice_service = InvoiceService(repository=invoice_repo)
    payment_service = PaymentService(
        payment_repository=payment_repo,
        invoice_repository=invoice_repo,
    )
    return invoice_service, payment_service, invoice_repo, payment_repo


def _serialize_invoice(invoice: Invoice) -> InvoiceResponse:
    return InvoiceResponse(
        id=invoice.id,
        citizen_id=invoice.citizen_id,
        reference=invoice.reference,
        revenue_code=invoice.revenue_code,
        cost_center=invoice.cost_center,
        service_code=invoice.service_code,
        service_name=invoice.service_name,
        amount=float(invoice.amount),
        currency=invoice.currency,
        status=invoice.status.value if hasattr(invoice.status, "value") else str(invoice.status),
        created_at=invoice.created_at.isoformat(),
        updated_at=invoice.updated_at.isoformat(),
        due_date=invoice.due_date.isoformat(),
    )


def _serialize_payment(payment: Payment) -> PaymentResponse:
    return PaymentResponse(
        id=payment.id,
        invoice_id=payment.invoice_id,
        citizen_id=payment.citizen_id,
        amount=float(payment.amount),
        currency=payment.currency,
        gateway_reference=payment.gateway_reference,
        payment_method=payment.payment_method,
        status=payment.status.value if hasattr(payment.status, "value") else str(payment.status),
        created_at=payment.created_at.isoformat(),
        confirmation_timestamp=payment.confirmed_at.isoformat() if payment.confirmed_at else None,
    )


@router.get("/")
def root() -> dict[str, str]:
    return {"message": "SILA Backend online"}


@router.get("/health")
def health() -> dict[str, Any]:
    return {
        "status": "healthy",
        "service": "SILA-System",
        "version": "20.2",
        "environment": "development",
        "timestamp": time.time(),
        "database": "local",
    }


@router.options("/health")
def health_options() -> Response:
    return Response(
        status_code=status.HTTP_200_OK,
        headers={
            "access-control-allow-origin": "*",
            "access-control-allow-methods": "GET,OPTIONS",
            "access-control-allow-headers": "*",
        },
    )


@router.get("/info")
def info() -> dict[str, Any]:
    return {
        "system": {"name": "SILA", "environment": "development", "debug": True, "version": "20.2"},
        "database": {"host": "localhost", "database": "sila_db"},
        "features": {"event_sourcing": True, "trust_engine": True, "catalog": True},
    }


@router.get("/auth/me")
@router.get("/api/auth/me")
@router.get("/api/v1/auth/me")
async def auth_me(current_user: dict[str, Any] = Depends(get_current_user)) -> dict[str, Any]:
    return current_user


@router.get("/api/auth/whoami")
async def auth_whoami(current_user: dict[str, Any] = Depends(get_current_user)) -> dict[str, Any]:
    return {
        "sub": current_user.get("sub"),
        "roles": current_user.get("roles", []),
        "scopes": current_user.get("scopes", []),
    }


def _normalize_roles(raw_roles: Any) -> list[str]:
    if raw_roles is None:
        return []
    if isinstance(raw_roles, list):
        return [str(role).upper() for role in raw_roles if role]
    if isinstance(raw_roles, str):
        import json

        try:
            parsed = json.loads(raw_roles)
            if isinstance(parsed, list):
                return [str(role).upper() for role in parsed if role]
        except Exception:
            pass
        return [raw_roles.upper()]
    return [str(raw_roles).upper()]


def _map_primary_role(roles: list[str], email: str | None) -> str:
    role_set = {role.upper() for role in roles if role}
    if "CITIZEN" in role_set:
        return "CITIZEN"
    if "SUPERADMIN" in role_set:
        return "ADMIN_SUPER"
    if "ADMIN" in role_set:
        return "ADMIN_CENTRAL"
    if "MANAGER" in role_set:
        prefix = (email or "").split("@")[0].lower()
        if prefix.startswith("prov"):
            return "ADMIN_PROVINCIAL"
        if prefix.startswith("mun"):
            return "ADMIN_MUNICIPAL"
        if prefix.startswith("comun"):
            return "ADMIN_COMMUNAL"
        return "ADMIN_CENTRAL"
    return "CITIZEN"


def _map_levels(primary_role: str) -> tuple[str | None, str | None]:
    if primary_role == "ADMIN_SUPER":
        return ("super", "CENTRAL")
    if primary_role == "ADMIN_CENTRAL":
        return ("central", "CENTRAL")
    if primary_role == "ADMIN_PROVINCIAL":
        return ("provincial", "PROVINCIAL")
    if primary_role == "ADMIN_MUNICIPAL":
        return ("municipal", "LOCAL")
    if primary_role == "ADMIN_COMMUNAL":
        return ("communal", "LOCAL")
    return (None, None)


async def _read_login_payload(request: Request) -> dict[str, Any]:
    content_type = request.headers.get("content-type", "")
    if "application/json" in content_type:
        try:
            return await request.json()
        except Exception:
            return {}
    if "application/x-www-form-urlencoded" in content_type or "multipart/form-data" in content_type:
        form = await request.form()
        return dict(form)
    try:
        return await request.json()
    except Exception:
        return {}


async def _fetch_iam_user(db: AsyncSession, identifier: str) -> dict[str, Any] | None:
    result = await db.execute(
        text(
            "\n            SELECT id, username, email, password_hash, status, is_active, is_superuser,\n                   full_name, phone, department, position, citizen_id, custom_metadata\n            FROM iam_users\n            WHERE lower(username) = lower(:identifier)\n               OR lower(email) = lower(:identifier)\n            LIMIT 1\n            "
        ),
        {"identifier": identifier},
    )
    row = result.mappings().first()
    return dict(row) if row else None


async def _fetch_iam_roles(db: AsyncSession, user_id: str) -> list[str]:
    result = await db.execute(
        text(
            "\n            SELECT r.name\n            FROM iam_roles r\n            JOIN iam_user_roles ur ON ur.role_id = r.id\n            WHERE ur.user_id = :user_id\n              AND (ur.is_active IS NULL OR ur.is_active = true)\n            "
        ),
        {"user_id": user_id},
    )
    return [row[0].upper() for row in result.fetchall()]


async def _fetch_legacy_user(db: AsyncSession, identifier: str) -> dict[str, Any] | None:
    result = await db.execute(
        text(
            "\n            SELECT *\n            FROM users\n            WHERE lower(username) = lower(:identifier)\n               OR lower(email) = lower(:identifier)\n            LIMIT 1\n            "
        ),
        {"identifier": identifier},
    )
    row = result.mappings().first()
    return dict(row) if row else None


async def _fetch_citizen_profile(
    db: AsyncSession, citizen_id: str | None, email: str | None
) -> dict[str, Any] | None:
    if citizen_id:
        result = await db.execute(
            text("""
                SELECT id, name, birth_date, bi_number, email, phone
                FROM citizenship_citizens
                WHERE id = :citizen_id
                LIMIT 1
            """),
            {"citizen_id": citizen_id},
        )
        row = result.mappings().first()
        if row:
            return dict(row)
    if email:
        result = await db.execute(
            text("""
                SELECT id, name, birth_date, bi_number, email, phone
                FROM citizenship_citizens
                WHERE lower(email) = lower(:email)
                LIMIT 1
            """),
            {"email": email},
        )
        row = result.mappings().first()
        if row:
            return dict(row)
    return None


async def _fetch_citizen_profile_from_iam(
    db: AsyncSession,
    email: str | None,
) -> dict[str, Any] | None:
    if not email:
        return None
    result = await db.execute(
        text(
            """
            SELECT full_name, email, phone, custom_metadata
            FROM iam_users
            WHERE lower(email) = lower(:email)
            LIMIT 1
            """
        ),
        {"email": email},
    )
    row = result.mappings().first()
    if not row:
        return None
    metadata = row.get("custom_metadata") or {}
    if isinstance(metadata, str):
        try:
            metadata = json.loads(metadata)
        except Exception:
            metadata = {}
    if not isinstance(metadata, dict):
        metadata = {}
    return {
        "id": metadata.get("citizen_id"),
        "name": metadata.get("full_name") or row.get("full_name") or row.get("email"),
        "birth_date": metadata.get("birth_date"),
        "bi_number": metadata.get("bi_number"),
        "email": row.get("email"),
        "phone": row.get("phone"),
    }


async def _fetch_citizen_events_from_projections(
    db: AsyncSession,
    citizen_id: str | None,
    email: str | None,
) -> list[dict[str, Any]] | None:
    # TODO: swap the citizen portal events to event_store_projections when projections are finalized.
    projection_name = os.getenv("CITIZEN_EVENTS_PROJECTION")
    aggregate_id = citizen_id or email
    if not projection_name or not aggregate_id:
        return None
    result = await db.execute(
        text("""
            SELECT data
            FROM event_store_projections
            WHERE projection_name = :projection_name
              AND aggregate_id = :aggregate_id
            LIMIT 1
        """),
        {"projection_name": projection_name, "aggregate_id": str(aggregate_id)},
    )
    row = result.mappings().first()
    if not row:
        return None
    data = row.get("data") or {}
    if not isinstance(data, dict):
        return None
    items = data.get("events")
    if not isinstance(items, list):
        return None
    events: list[dict[str, Any]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        events.append(
            {
                "id": str(item.get("id") or uuid.uuid4()),
                "event_type": item.get("type") or item.get("event_type") or "Evento",
                "created_at": item.get("timestamp") or item.get("created_at"),
                "payload": item.get("payload") or {},
            }
        )
    return events


def _serialize_datetime(value: Any) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _parse_is_active(status_value: str | None) -> bool | None:
    if not status_value:
        return None
    normalized = status_value.strip().lower()
    if normalized in {"active", "ativo", "true", "1"}:
        return True
    if normalized in {"inactive", "inativo", "false", "0"}:
        return False
    return None


def _parse_datetime(value: str | None, end: bool = False) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
        return parsed
    except ValueError:
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            return parsed
        except Exception:
            return None


def _assert_admin(current_user: dict[str, Any]) -> None:
    role = (current_user.get("role") or "").upper()
    if not role or role == "CITIZEN":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")


async def _safe_scalar(
    db: AsyncSession,
    query: str,
    params: dict[str, Any] | None = None,
) -> tuple[Any | None, bool]:
    try:
        result = await db.execute(text(query), params or {})
        return result.scalar_one_or_none(), True
    except Exception:
        return None, False


async def _safe_count(
    db: AsyncSession,
    query: str,
    params: dict[str, Any] | None = None,
) -> tuple[int, bool]:
    value, ok = await _safe_scalar(db, query, params)
    if value is None:
        return 0, ok
    try:
        return int(value), ok
    except Exception:
        return 0, False


def _sum_metrics(prefix: str) -> int:
    total = 0
    for key, value in Metrics._metrics.items():
        if not str(key).startswith(prefix):
            continue
        try:
            total += int(value)
        except Exception:
            continue
    return total


def _map_service_type_to_sla_code(service_type: str | None) -> str:
    if not service_type:
        return "DEFAULT"
    mapping = {
        "identity_bi": "BI_EMISSAO",
        "civil_birth": "CERTIDAO_NASCIMENTO",
        "civil_marriage": "CERTIDAO_CASAMENTO",
        "civil_death": "CERTIDAO_OBITO",
    }
    normalized = service_type.strip().lower()
    upper = normalized.upper()
    if upper in SLA_DEFINITIONS:
        return upper
    return mapping.get(normalized, "DEFAULT")


def _map_citizen_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(row.get("id")),
        "full_name": row.get("name"),
        "bi_number": row.get("bi_number"),
        "birth_date": _serialize_datetime(row.get("birth_date")),
        "email": row.get("email"),
        "phone": row.get("phone"),
        "is_active": row.get("is_active"),
        "status": "ACTIVE" if row.get("is_active") else "INACTIVE",
        "created_at": _serialize_datetime(row.get("created_at")),
        "updated_at": _serialize_datetime(row.get("updated_at")),
        "address": row.get("address"),
        "birth_location_id": row.get("birth_location_id"),
        "residence_location_id": row.get("residence_location_id"),
        "user_id": row.get("user_id"),
    }


def _document_status(valid_until: Any) -> str:
    if not valid_until:
        return "UNKNOWN"
    if isinstance(valid_until, datetime):
        return "EXPIRED" if valid_until < datetime.utcnow() else "ACTIVE"
    return "UNKNOWN"


def _map_document_row(row: dict[str, Any]) -> dict[str, Any]:
    status = _document_status(row.get("valid_until"))
    return {
        "id": str(row.get("id")),
        "citizen_id": row.get("citizen_id"),
        "citizen_name": row.get("citizen_name"),
        "citizen_bi": row.get("bi_number"),
        "citizen_email": row.get("citizen_email"),
        "document_type": row.get("document_type"),
        "file_url": row.get("file_url"),
        "issued_at": _serialize_datetime(row.get("issued_at")),
        "valid_until": _serialize_datetime(row.get("valid_until")),
        "status": status,
    }


def _map_territory_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": str(row.get("id")),
        "name": row.get("name"),
        "code": row.get("code"),
        "type": row.get("type"),
        "parent_id": str(row.get("parent_id")) if row.get("parent_id") else None,
    }


async def _build_citizens_export_bytes(
    db: AsyncSession,
    params: dict[str, Any],
    columns: str | None,
    export_format: str,
) -> tuple[bytes, str, str]:
    result = await db.execute(
        text("""
            SELECT id, name, email, phone, address, is_active, created_at, updated_at,
                   bi_number, birth_date, birth_location_id, residence_location_id, user_id
            FROM citizenship_citizens
            WHERE (:name IS NULL OR name ILIKE :name_like)
              AND (:bi IS NULL OR bi_number ILIKE :bi_like)
              AND (:email IS NULL OR email ILIKE :email_like)
              AND (:phone IS NULL OR phone ILIKE :phone_like)
              AND (:is_active IS NULL OR is_active = :is_active)
              AND (:created_from IS NULL OR created_at >= :created_from)
              AND (:created_to IS NULL OR created_at <= :created_to)
              AND (
                :q IS NULL
                OR name ILIKE :q_like
                OR email ILIKE :q_like
                OR bi_number ILIKE :q_like
                OR phone ILIKE :q_like
              )
            ORDER BY created_at DESC NULLS LAST
        """),
        params,
    )
    rows = [_map_citizen_row(dict(row)) for row in result.mappings().all()]
    allowed_columns = [
        "id",
        "full_name",
        "bi_number",
        "birth_date",
        "email",
        "phone",
        "is_active",
        "status",
        "created_at",
        "updated_at",
    ]
    default_columns = [
        "full_name",
        "bi_number",
        "birth_date",
        "email",
        "phone",
        "status",
        "created_at",
    ]
    fieldnames = _parse_columns(columns, allowed_columns, default_columns)
    if export_format == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Citizens"
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(key) for key in fieldnames])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"citizens_{datetime.utcnow().date().isoformat()}.xlsx"
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename,
        )
    import csv
    import io

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) for key in fieldnames})
    csv_content = "\ufeff" + output.getvalue()
    filename = f"citizens_{datetime.utcnow().date().isoformat()}.csv"
    return csv_content.encode("utf-8"), "text/csv", filename


async def _build_documents_export_bytes(
    db: AsyncSession,
    params: dict[str, Any],
    columns: str | None,
    export_format: str,
) -> tuple[bytes, str, str]:
    result = await db.execute(
        text("""
            SELECT d.id,
                   d.citizen_id,
                   d.document_type,
                   d.file_url,
                   d.issued_at,
                   d.valid_until,
                   c.name AS citizen_name,
                   c.bi_number AS bi_number,
                   c.email AS citizen_email
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE (:document_type IS NULL OR d.document_type ILIKE :document_type_like)
              AND (:bi IS NULL OR c.bi_number ILIKE :bi_like)
              AND (:email IS NULL OR c.email ILIKE :email_like)
              AND (:issued_from IS NULL OR d.issued_at >= :issued_from)
              AND (:issued_to IS NULL OR d.issued_at <= :issued_to)
              AND (
                :status IS NULL
                OR (:status = 'active' AND (d.valid_until IS NULL OR d.valid_until >= NOW()))
                OR (:status = 'expired' AND d.valid_until < NOW())
                OR (:status = 'unknown' AND d.valid_until IS NULL)
              )
              AND (
                :q IS NULL
                OR d.document_type ILIKE :q_like
                OR c.name ILIKE :q_like
                OR c.bi_number ILIKE :q_like
                OR c.email ILIKE :q_like
              )
            ORDER BY d.issued_at DESC NULLS LAST
        """),
        params,
    )
    rows = [_map_document_row(dict(row)) for row in result.mappings().all()]
    allowed_columns = [
        "id",
        "document_type",
        "citizen_name",
        "citizen_bi",
        "citizen_email",
        "issued_at",
        "valid_until",
        "status",
        "file_url",
        "citizen_id",
    ]
    default_columns = [
        "document_type",
        "citizen_name",
        "citizen_bi",
        "issued_at",
        "valid_until",
        "status",
        "file_url",
    ]
    fieldnames = _parse_columns(columns, allowed_columns, default_columns)
    if export_format == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Documents"
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(key) for key in fieldnames])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"documents_{datetime.utcnow().date().isoformat()}.xlsx"
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename,
        )
    import csv
    import io

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) for key in fieldnames})
    csv_content = "\ufeff" + output.getvalue()
    filename = f"documents_{datetime.utcnow().date().isoformat()}.csv"
    return csv_content.encode("utf-8"), "text/csv", filename


async def _count_citizens_export(db: AsyncSession, params: dict[str, Any]) -> int:
    result = await db.execute(
        text("""
            SELECT count(*)
            FROM citizenship_citizens
            WHERE (:name IS NULL OR name ILIKE :name_like)
              AND (:bi IS NULL OR bi_number ILIKE :bi_like)
              AND (:email IS NULL OR email ILIKE :email_like)
              AND (:phone IS NULL OR phone ILIKE :phone_like)
              AND (:is_active IS NULL OR is_active = :is_active)
              AND (:created_from IS NULL OR created_at >= :created_from)
              AND (:created_to IS NULL OR created_at <= :created_to)
              AND (
                :q IS NULL
                OR name ILIKE :q_like
                OR email ILIKE :q_like
                OR bi_number ILIKE :q_like
                OR phone ILIKE :q_like
              )
        """),
        params,
    )
    return int(result.scalar_one() or 0)


async def _count_documents_export(db: AsyncSession, params: dict[str, Any]) -> int:
    result = await db.execute(
        text("""
            SELECT count(*)
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE (:document_type IS NULL OR d.document_type ILIKE :document_type_like)
              AND (:bi IS NULL OR c.bi_number ILIKE :bi_like)
              AND (:email IS NULL OR c.email ILIKE :email_like)
              AND (:issued_from IS NULL OR d.issued_at >= :issued_from)
              AND (:issued_to IS NULL OR d.issued_at <= :issued_to)
              AND (
                :status IS NULL
                OR (:status = 'active' AND (d.valid_until IS NULL OR d.valid_until >= NOW()))
                OR (:status = 'expired' AND d.valid_until < NOW())
                OR (:status = 'unknown' AND d.valid_until IS NULL)
              )
              AND (
                :q IS NULL
                OR d.document_type ILIKE :q_like
                OR c.name ILIKE :q_like
                OR c.bi_number ILIKE :q_like
                OR c.email ILIKE :q_like
              )
        """),
        params,
    )
    return int(result.scalar_one() or 0)


async def _count_documents_by_status(db: AsyncSession, params: dict[str, Any]) -> dict[str, int]:
    base_params = {**params}
    base_params["status"] = None
    counts = {
        "active": await _count_documents_export(db, {**base_params, "status": "active"}),
        "expired": await _count_documents_export(db, {**base_params, "status": "expired"}),
        "unknown": await _count_documents_export(db, {**base_params, "status": "unknown"}),
    }
    return counts


async def _count_citizens_by_status(db: AsyncSession, params: dict[str, Any]) -> dict[str, int]:
    base_params = {**params}
    base_params["is_active"] = None
    counts = {
        "active": await _count_citizens_export(db, {**base_params, "is_active": True}),
        "inactive": await _count_citizens_export(db, {**base_params, "is_active": False}),
    }
    return counts


def _parse_date(value: str | None, end: bool = False) -> datetime | None:
    if not value:
        return None
    try:
        if "T" in value:
            return datetime.fromisoformat(value)
        parsed = datetime.fromisoformat(value).date()
        return datetime.combine(parsed, time.max if end else time.min)
    except Exception:
        return None


def _parse_columns(columns: str | None, allowed: list[str], default: list[str]) -> list[str]:
    if not columns:
        return default
    raw = [item.strip() for item in columns.split(",") if item.strip()]
    if not raw:
        return default
    filtered = [item for item in raw if item in allowed]
    return filtered or default


async def _fetch_iam_user_record(
    db: AsyncSession, user_id: str | None, email: str | None
) -> dict[str, Any] | None:
    if user_id:
        result = await db.execute(
            text("""
                SELECT id, email, custom_metadata
                FROM iam_users
                WHERE id = :user_id
                LIMIT 1
            """),
            {"user_id": user_id},
        )
        row = result.mappings().first()
        if row:
            return dict(row)
    if email:
        result = await db.execute(
            text("""
                SELECT id, email, custom_metadata
                FROM iam_users
                WHERE lower(email) = lower(:email)
                LIMIT 1
            """),
            {"email": email},
        )
        row = result.mappings().first()
        if row:
            return dict(row)
    return None


def _normalize_metadata(value: Any) -> dict[str, Any]:
    if value is None:
        return {}
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            return {}
    return {}


async def _load_export_presets(
    db: AsyncSession, current_user: dict[str, Any], module: str
) -> list[dict[str, Any]]:
    user_id = current_user.get("id") or current_user.get("sub")
    email = current_user.get("email")
    record = await _fetch_iam_user_record(db, user_id, email)
    if not record:
        return []
    metadata = _normalize_metadata(record.get("custom_metadata"))
    presets = metadata.get("export_presets", {})
    module_presets = presets.get(module, {})
    if isinstance(module_presets, dict):
        return [{"name": name, "columns": columns} for name, columns in module_presets.items()]
    return []


async def _save_export_preset(
    db: AsyncSession, current_user: dict[str, Any], module: str, name: str, columns: list[str]
) -> list[dict[str, Any]]:
    user_id = current_user.get("id") or current_user.get("sub")
    email = current_user.get("email")
    record = await _fetch_iam_user_record(db, user_id, email)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    metadata = _normalize_metadata(record.get("custom_metadata"))
    presets = metadata.setdefault("export_presets", {})
    module_presets = presets.setdefault(module, {})
    if not isinstance(module_presets, dict):
        module_presets = {}
        presets[module] = module_presets
    module_presets[name] = columns
    await db.execute(
        text("""
            UPDATE iam_users
            SET custom_metadata = :metadata::jsonb
            WHERE id = :user_id
        """),
        {"metadata": json.dumps(metadata), "user_id": record["id"]},
    )
    await db.commit()
    return [
        {"name": preset_name, "columns": preset_cols}
        for preset_name, preset_cols in module_presets.items()
    ]


async def _delete_export_preset(
    db: AsyncSession, current_user: dict[str, Any], module: str, name: str
) -> list[dict[str, Any]]:
    user_id = current_user.get("id") or current_user.get("sub")
    email = current_user.get("email")
    record = await _fetch_iam_user_record(db, user_id, email)
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")
    metadata = _normalize_metadata(record.get("custom_metadata"))
    presets = metadata.get("export_presets", {})
    module_presets = presets.get(module, {})
    if isinstance(module_presets, dict) and name in module_presets:
        module_presets.pop(name, None)
    await db.execute(
        text("""
            UPDATE iam_users
            SET custom_metadata = :metadata::jsonb
            WHERE id = :user_id
        """),
        {"metadata": json.dumps(metadata), "user_id": record["id"]},
    )
    await db.commit()
    if isinstance(module_presets, dict):
        return [
            {"name": preset_name, "columns": preset_cols}
            for preset_name, preset_cols in module_presets.items()
        ]
    return []


async def _create_export_job(
    db: AsyncSession,
    module: str,
    owner_id: str,
    payload: dict[str, Any],
) -> str:
    result = await db.execute(
        text("""
            INSERT INTO export_jobs (id, module, owner_id, status, request_payload, created_at, updated_at)
            VALUES (gen_random_uuid(), :module, :owner_id, 'pending', :payload::jsonb, now(), now())
            RETURNING id
        """),
        {"module": module, "owner_id": owner_id, "payload": json.dumps(payload)},
    )
    job_id = str(result.scalar_one())
    await db.commit()
    try:
        await db.execute(
            text("""
                INSERT INTO export_job_logs (id, job_id, level, message, created_at)
                VALUES (gen_random_uuid(), :job_id, 'info', :message, now())
            """),
            {"job_id": job_id, "message": "Job queued"},
        )
        await db.commit()
    except Exception:
        await db.rollback()
    return job_id


async def _fetch_export_job(db: AsyncSession, job_id: str) -> dict[str, Any] | None:
    result = await db.execute(
        text("""
            SELECT id, module, owner_id, status, request_payload,
                   result_data, result_content_type, result_filename,
                   error, created_at, updated_at
            FROM export_jobs
            WHERE id = :job_id
            LIMIT 1
        """),
        {"job_id": job_id},
    )
    row = result.mappings().first()
    return dict(row) if row else None


async def _update_export_job(
    db: AsyncSession,
    job_id: str,
    status_value: str,
    result_data: bytes | None = None,
    result_content_type: str | None = None,
    result_filename: str | None = None,
    error: str | None = None,
) -> None:
    await db.execute(
        text("""
            UPDATE export_jobs
            SET status = :status,
                result_data = COALESCE(:result_data, result_data),
                result_content_type = COALESCE(:result_content_type, result_content_type),
                result_filename = COALESCE(:result_filename, result_filename),
                error = COALESCE(:error, error),
                updated_at = now()
            WHERE id = :job_id
        """),
        {
            "status": status_value,
            "result_data": result_data,
            "result_content_type": result_content_type,
            "result_filename": result_filename,
            "error": error,
            "job_id": job_id,
        },
    )
    await db.commit()


async def _list_export_jobs(
    db: AsyncSession,
    owner_id: str,
    module: str | None,
    status_value: str | None,
    query: str | None,
    job_id: str | None,
    date_from: datetime | None,
    date_to: datetime | None,
    limit: int,
    offset: int,
) -> tuple[list[dict[str, Any]], int, dict[str, float], dict[str, int]]:
    q_like = f"%{query}%" if query else None
    result = await db.execute(
        text("""
            SELECT count(*)
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND (CAST(:job_id AS TEXT) IS NULL OR id::text = CAST(:job_id AS TEXT))
              AND (CAST(:module AS TEXT) IS NULL OR module = CAST(:module AS TEXT))
              AND (CAST(:status AS TEXT) IS NULL OR status = CAST(:status AS TEXT))
              AND (CAST(:date_from AS TIMESTAMPTZ) IS NULL OR created_at >= CAST(:date_from AS TIMESTAMPTZ))
              AND (CAST(:date_to AS TIMESTAMPTZ) IS NULL OR created_at <= CAST(:date_to AS TIMESTAMPTZ))
              AND (
                    CAST(:q AS TEXT) IS NULL
                    OR module ILIKE CAST(:q_like AS TEXT)
                    OR status ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(result_filename, '') ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(error, '') ILIKE CAST(:q_like AS TEXT)
                  )
        """),
        {
            "owner_id": owner_id,
            "job_id": job_id,
            "module": module,
            "status": status_value,
            "date_from": date_from,
            "date_to": date_to,
            "q": query,
            "q_like": q_like,
        },
    )
    total = int(result.scalar_one() or 0)
    avg_rows = await db.execute(
        text("""
            SELECT module,
                   AVG(EXTRACT(EPOCH FROM (updated_at - created_at))) AS avg_seconds
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND status = 'done'
            GROUP BY module
        """),
        {"owner_id": owner_id},
    )
    avg_map = {row["module"]: float(row["avg_seconds"] or 0.0) for row in avg_rows.mappings().all()}
    status_rows = await db.execute(
        text("""
            SELECT status, count(*) AS count
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND (CAST(:job_id AS TEXT) IS NULL OR id::text = CAST(:job_id AS TEXT))
              AND (CAST(:module AS TEXT) IS NULL OR module = CAST(:module AS TEXT))
              AND (CAST(:date_from AS TIMESTAMPTZ) IS NULL OR created_at >= CAST(:date_from AS TIMESTAMPTZ))
              AND (CAST(:date_to AS TIMESTAMPTZ) IS NULL OR created_at <= CAST(:date_to AS TIMESTAMPTZ))
              AND (
                    CAST(:q AS TEXT) IS NULL
                    OR module ILIKE CAST(:q_like AS TEXT)
                    OR status ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(result_filename, '') ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(error, '') ILIKE CAST(:q_like AS TEXT)
                  )
            GROUP BY status
        """),
        {
            "owner_id": owner_id,
            "job_id": job_id,
            "module": module,
            "date_from": date_from,
            "date_to": date_to,
            "q": query,
            "q_like": q_like,
        },
    )
    status_counts = {row["status"]: int(row["count"] or 0) for row in status_rows.mappings().all()}
    rows = await db.execute(
        text("""
            SELECT id AS job_id,
                   module,
                   status,
                   result_filename,
                   error,
                   created_at,
                   updated_at
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND (CAST(:job_id AS TEXT) IS NULL OR id::text = CAST(:job_id AS TEXT))
              AND (CAST(:module AS TEXT) IS NULL OR module = CAST(:module AS TEXT))
              AND (CAST(:status AS TEXT) IS NULL OR status = CAST(:status AS TEXT))
              AND (CAST(:date_from AS TIMESTAMPTZ) IS NULL OR created_at >= CAST(:date_from AS TIMESTAMPTZ))
              AND (CAST(:date_to AS TIMESTAMPTZ) IS NULL OR created_at <= CAST(:date_to AS TIMESTAMPTZ))
              AND (
                    CAST(:q AS TEXT) IS NULL
                    OR module ILIKE CAST(:q_like AS TEXT)
                    OR status ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(result_filename, '') ILIKE CAST(:q_like AS TEXT)
                    OR COALESCE(error, '') ILIKE CAST(:q_like AS TEXT)
                  )
            ORDER BY created_at DESC
            LIMIT :limit OFFSET :offset
        """),
        {
            "owner_id": owner_id,
            "job_id": job_id,
            "module": module,
            "status": status_value,
            "date_from": date_from,
            "date_to": date_to,
            "q": query,
            "q_like": q_like,
            "limit": limit,
            "offset": offset,
        },
    )
    items = []
    for row in rows.mappings().all():
        item = dict(row)
        duration_seconds: float | None = None
        if item.get("created_at") and item.get("updated_at"):
            duration_seconds = (item["updated_at"] - item["created_at"]).total_seconds()
        if item.get("created_at"):
            item["created_at"] = _serialize_datetime(item["created_at"])
        if item.get("updated_at"):
            item["updated_at"] = _serialize_datetime(item["updated_at"])
        if duration_seconds is not None:
            item["duration_seconds"] = round(duration_seconds, 2)
        else:
            item["duration_seconds"] = None
        avg_seconds = avg_map.get(item.get("module") or "", 0.0)
        if item.get("status") in {"pending", "running"} and avg_seconds > 0:
            item["estimated_seconds"] = round(avg_seconds, 2)
        else:
            item["estimated_seconds"] = None
        items.append(item)
    return items, total, avg_map, status_counts


def _export_job_response(job_id: str, job: dict[str, Any]) -> dict[str, Any]:
    response = {
        "job_id": job_id,
        "status": job.get("status"),
    }
    if job.get("status") == "done":
        response["download_url"] = f"/api/admin/exports/jobs/{job_id}/download"
        response["filename"] = job.get("result_filename")
    if job.get("status") == "failed":
        response["error"] = job.get("error")
    return response


def _extract_bearer_token(raw_value: str | None) -> str | None:
    if not raw_value:
        return None
    value = raw_value.strip()
    if value.lower().startswith("bearer "):
        return value.split(" ", 1)[1].strip()
    return value


def _decode_token(token: str) -> dict[str, Any]:
    jwt_handler = JWTHandler(secret_key=settings.SECRET_KEY)
    return jwt_handler.decode_token(token)


async def _resolve_stream_user(
    request: Request,
    token: str | None,
    authorization: str | None,
) -> dict[str, Any]:
    raw_token = token or _extract_bearer_token(authorization)
    if not raw_token:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Missing token")
    try:
        return _decode_token(raw_token)
    except Exception:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")


def _build_user_payload(user: dict[str, Any], roles: list[str], source: str) -> dict[str, Any]:
    primary_role = _map_primary_role(roles, user.get("email"))
    level_lower, level_upper = _map_levels(primary_role)
    payload: dict[str, Any] = {
        "sub": str(user.get("id")),
        "id": str(user.get("id")),
        "email": user.get("email"),
        "username": user.get("username") or user.get("email"),
        "full_name": user.get("full_name"),
        "roles": roles,
        "role": primary_role,
        "level": level_lower,
        "administrative_level": level_upper,
        "is_active": user.get("is_active", True),
        "status": user.get("status") or "ACTIVE",
        "source": source,
    }
    if user.get("citizen_id"):
        payload["citizen_id"] = str(user.get("citizen_id"))
    return payload


def _issue_access_token(claims: dict[str, Any]) -> str:
    expires = timedelta(minutes=settings.ACCESS_TOKEN_EXPIRE_MINUTES)
    jwt_handler = JWTHandler(secret_key=settings.SECRET_KEY)
    full_claims = {**claims, "type": "access"}
    # Use the internal _encode_token method to pass custom claims
    return jwt_handler._encode_token(full_claims, expires)


def _issue_refresh_token(claims: dict[str, Any]) -> str:
    expires = timedelta(days=max(1, settings.ACCESS_TOKEN_EXPIRE_DAYS * 2))
    jwt_handler = JWTHandler(secret_key=settings.SECRET_KEY)
    refresh_claims = {**claims, "type": "refresh"}
    # Use the internal _encode_token method to pass custom claims
    return jwt_handler._encode_token(refresh_claims, expires)


@router.post("/auth/login")
@router.post("/api/auth/login")
@router.post("/api/v1/auth/login")
@router.post("/auth/login/access-token")
@router.post("/api/v1/auth/login/access-token")
async def auth_login(request: Request, db: AsyncSession = Depends(get_db)) -> dict[str, Any]:
    payload = await _read_login_payload(request)
    identifier = (payload.get("username") or payload.get("email") or "").strip()
    password = (payload.get("password") or "").strip()
    if not identifier or not password:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing credentials")
    user: dict[str, Any] | None = None
    roles: list[str] = []
    source = "iam_users"
    try:
        user = await _fetch_iam_user(db, identifier)
        if user:
            roles = await _fetch_iam_roles(db, str(user.get("id")))
            if not roles and user.get("is_superuser"):
                roles = ["SUPERADMIN"]
    except Exception:
        await db.rollback()
        user = None
    if not user and os.getenv("LEGACY_USERS_FALLBACK", "0") == "1":
        source = "users"
        try:
            legacy_user = await _fetch_legacy_user(db, identifier)
        except Exception:
            await db.rollback()
            legacy_user = None
        if legacy_user:
            user = legacy_user
            roles = _normalize_roles(user.get("roles"))
            if not roles and user.get("is_superuser"):
                roles = ["SUPERADMIN"]
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Credenciais inválidas"
        )
    password_hash = user.get("password_hash") or user.get("hashed_password")
    if not password_hash or not verify_password(password, password_hash):
        try:
            await db.execute(
                text(
                    "\n                    UPDATE iam_users\n                    SET failed_login_attempts = failed_login_attempts + 1\n                    WHERE id = :user_id\n                    "
                ),
                {"user_id": user.get("id")},
            )
            await db.commit()
        except Exception:
            pass
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Credenciais inválidas"
        )
    if not user.get("is_active", True):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Conta inativa")
    try:
        await db.execute(
            text(
                "\n                UPDATE iam_users\n                SET last_login_at = :now,\n                    last_login_ip = :ip,\n                    failed_login_attempts = 0\n                WHERE id = :user_id\n                "
            ),
            {
                "now": datetime.utcnow(),
                "ip": request.client.host if request.client else None,
                "user_id": user.get("id"),
            },
        )
        await db.commit()
    except Exception:
        pass
    claims = _build_user_payload(user, roles, source)
    access_token = _issue_access_token(claims)
    refresh_token = _issue_refresh_token(claims)
    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "token_type": "bearer",
        "user": claims,
    }


@router.get("/citizen/profile")
@router.get("/api/citizen/profile")
@router.get("/api/v1/citizen/profile")
async def citizen_profile(
    current_user: dict[str, Any] = Depends(get_current_user), db: AsyncSession = Depends(get_db)
) -> dict[str, Any]:
    citizen_id = current_user.get("citizen_id")
    email = current_user.get("email")
    profile = await _fetch_citizen_profile(db, citizen_id, email)
    if not profile:
        profile = await _fetch_citizen_profile_from_iam(db, email)
    if profile:
        return {
            "id": str(profile.get("id") or citizen_id or current_user.get("id") or ""),
            "full_name": profile.get("name") or current_user.get("full_name") or email,
            "birth_date": profile.get("birth_date"),
            "gender": None,
            "vital_status": None,
            "id_number": profile.get("bi_number"),
            "nif": None,
            "email": profile.get("email") or email,
            "phone": profile.get("phone"),
        }
    return {
        "id": str(citizen_id or current_user.get("id") or ""),
        "full_name": current_user.get("full_name") or email,
        "birth_date": None,
        "gender": None,
        "vital_status": None,
        "id_number": None,
        "nif": None,
        "email": email,
        "phone": None,
    }


@router.get("/citizen/events")
@router.get("/api/citizen/events")
@router.get("/api/v1/citizen/events")
async def citizen_events(
    current_user: dict[str, Any] = Depends(get_current_user), db: AsyncSession = Depends(get_db)
) -> list[dict[str, Any]]:
    citizen_id = current_user.get("citizen_id")
    email = current_user.get("email")
    events: list[dict[str, Any]] = []
    projections_events = await _fetch_citizen_events_from_projections(db, citizen_id, email)
    if projections_events is not None:
        return projections_events
    params = {
        "citizen_id": str(citizen_id) if citizen_id is not None else None,
        "email": str(email) if email is not None else None,
    }
    query = text("""
            SELECT n.id AS notification_id,
                   n.title AS title,
                   n.message AS message,
                   nd.delivered_at AS delivered_at
            FROM notifications n
            LEFT JOIN notification_deliveries nd
              ON nd.notification_id = n.id
            WHERE (:citizen_id IS NOT NULL AND n.citizen_id = :citizen_id)
               OR (:citizen_id IS NULL AND :email IS NOT NULL AND lower(n.citizen_id) = lower(:email))
            ORDER BY nd.delivered_at DESC NULLS LAST, n.title
            LIMIT 50
        """).bindparams(
        bindparam("citizen_id", type_=String),
        bindparam("email", type_=String),
    )
    result = await db.execute(query, params)
    for row in result.mappings().all():
        events.append(
            {
                "id": str(row.get("notification_id")),
                "event_type": row.get("title") or "Notificação",
                "created_at": row.get("delivered_at"),
                "payload": {"message": row.get("message")},
            }
        )
    return events


@router.get("/api/citizen/notifications")
@router.get("/citizen/notifications")
async def citizen_notifications(
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[dict[str, Any]]:
    return []


@router.get("/api/citizen/notifications/unread-count")
@router.get("/citizen/notifications/unread-count")
async def citizen_notifications_unread(
    current_user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, int]:
    return {"unread_count": 0}


@router.post("/api/citizen/notifications/{notification_id}/read")
@router.post("/citizen/notifications/{notification_id}/read")
async def citizen_notifications_mark_read(
    notification_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    return {"status": "ok", "id": notification_id}


@router.post("/api/citizen/notifications/read-all")
@router.post("/citizen/notifications/read-all")
async def citizen_notifications_mark_all(
    current_user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    return {"status": "ok"}


@router.get("/api/admin/dashboard")
@router.get("/api/admin/dashboard/")
async def admin_dashboard(
    current_user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    return {
        "metrics": {
            "total_citizens": "—",
            "total_documents": "—",
            "pending_requests": "—",
            "avg_response_time": "—",
            "trends": {
                "citizens": "+0%",
                "documents": "+0%",
                "requests": "+0%",
                "response_time": "+0%",
            },
        },
        "recent_events": [],
        "period": {"start": datetime.utcnow().isoformat(), "end": datetime.utcnow().isoformat()},
    }


@router.get("/api/admin/requests")
@router.get("/api/admin/requests/")
async def admin_requests(
    limit: int = 10,
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[dict[str, Any]]:
    return []


@router.get("/api/admin/observability/kpis")
async def admin_observability_kpis(
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    users_count, users_ok = await _safe_count(
        db,
        "SELECT COUNT(*) FROM iam_users",
    )
    if not users_ok:
        users_count, users_ok = await _safe_count(
            db,
            "SELECT COUNT(*) FROM users",
        )

    territories_count, territories_ok = await _safe_count(
        db,
        "SELECT COUNT(*) FROM locations",
    )
    services_count, services_ok = await _safe_count(
        db,
        "SELECT COUNT(DISTINCT service_type) FROM service_requests",
    )
    requests_count, requests_ok = await _safe_count(
        db,
        "SELECT COUNT(*) FROM service_requests",
    )

    audit_since = datetime.utcnow() - timedelta(hours=24)
    audit_errors_24h, audit_ok = await _safe_count(
        db,
        """
        SELECT COUNT(*)
        FROM audit_logs
        WHERE severity IN ('ERROR', 'CRITICAL')
          AND timestamp >= :since
        """,
        {"since": audit_since},
    )
    if not audit_ok:
        audit_errors_24h, audit_ok = await _safe_count(
            db,
            """
            SELECT COUNT(*)
            FROM audit_logs
            WHERE severity IN ('ERROR', 'CRITICAL')
              AND created_at >= :since
            """,
            {"since": audit_since},
        )

    in_memory_requests = _sum_metrics("request.total")
    in_memory_errors = _sum_metrics("request.error")
    traces_approx = in_memory_requests if OTEL_AVAILABLE else 0

    health_ok = all([users_ok, territories_ok, services_ok, requests_ok, audit_ok])
    scope = current_user.get("level") or current_user.get("administrative_level") or "central"
    target_id = current_user.get("territory_id") or current_user.get("location_id")

    return {
        "kpis": {
            "users": users_count,
            "territories": territories_count,
            "services": services_count,
            "requests": requests_count,
            "errors_24h": audit_errors_24h,
            "traces": traces_approx,
            "otel_metrics": 1 if OTEL_AVAILABLE else 0,
        },
        "meta": {
            "scope": scope,
            "target_id": str(target_id) if target_id else None,
            "health": "ok" if health_ok else "degraded",
        },
        "telemetry": {
            "otel_available": OTEL_AVAILABLE,
            "otel_tracing": OTEL_AVAILABLE,
            "otel_metrics": OTEL_AVAILABLE,
            "in_memory_requests": in_memory_requests,
            "in_memory_errors": in_memory_errors,
            "audit_errors_24h": audit_errors_24h,
        },
    }


@router.get("/api/admin/audit/sla/metrics")
async def admin_audit_sla_metrics(
    days: int = Query(7, ge=1, le=365),
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    since = datetime.utcnow() - timedelta(days=days)
    avg_seconds: float | None = None
    min_seconds: float | None = None
    max_seconds: float | None = None
    total_processed = 0
    try:
        result = await db.execute(
            text("""
                SELECT
                    COUNT(*) AS total_processed,
                    AVG(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS avg_seconds,
                    MIN(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS min_seconds,
                    MAX(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS max_seconds
                FROM service_requests
                WHERE completed_at IS NOT NULL
                  AND completed_at >= :since
            """),
            {"since": since},
        )
        row = result.mappings().first() or {}
        total_processed = int(row.get("total_processed") or 0)
        avg_seconds = float(row.get("avg_seconds")) if row.get("avg_seconds") is not None else None
        min_seconds = float(row.get("min_seconds")) if row.get("min_seconds") is not None else None
        max_seconds = float(row.get("max_seconds")) if row.get("max_seconds") is not None else None
    except Exception:
        total_processed = 0

    target_seconds = SLA_DEFINITIONS.get("DEFAULT", {}).get("target_sla_seconds", 259200)
    sla_status = evaluate_sla_status(avg_seconds, "DEFAULT") if avg_seconds is not None else None

    return {
        "avg_issuance_time_seconds": avg_seconds,
        "min_issuance_time_seconds": min_seconds,
        "max_issuance_time_seconds": max_seconds,
        "total_processed": total_processed,
        "period_days": days,
        "sla_status": sla_status,
        "sla_target_seconds": target_seconds,
    }


@router.get("/api/admin/audit/sla/breakdown")
async def admin_audit_sla_breakdown(
    days: int = Query(30, ge=1, le=365),
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    since = datetime.utcnow() - timedelta(days=days)
    rows: list[dict[str, Any]] = []
    try:
        result = await db.execute(
            text("""
                SELECT service_type,
                       COUNT(*) AS total_processed,
                       AVG(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS avg_seconds,
                       MIN(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS min_seconds,
                       MAX(EXTRACT(EPOCH FROM (completed_at - COALESCE(submitted_at, created_at)))) AS max_seconds
                FROM service_requests
                WHERE completed_at IS NOT NULL
                  AND completed_at >= :since
                GROUP BY service_type
                ORDER BY service_type
            """),
            {"since": since},
        )
        rows = [dict(row) for row in result.mappings().all()]
    except Exception:
        rows = []

    items: list[dict[str, Any]] = []
    for row in rows:
        service_type = row.get("service_type")
        sla_code = _map_service_type_to_sla_code(service_type)
        sla_def = get_sla_for_service(sla_code)
        avg_seconds = float(row.get("avg_seconds")) if row.get("avg_seconds") is not None else None
        min_seconds = float(row.get("min_seconds")) if row.get("min_seconds") is not None else None
        max_seconds = float(row.get("max_seconds")) if row.get("max_seconds") is not None else None
        status = evaluate_sla_status(avg_seconds, sla_code) if avg_seconds is not None else None
        items.append(
            {
                "service_type": service_type,
                "sla_code": sla_code,
                "sla_description": sla_def.get("description"),
                "sla_target_seconds": sla_def.get("target_sla_seconds"),
                "avg_issuance_time_seconds": avg_seconds,
                "min_issuance_time_seconds": min_seconds,
                "max_issuance_time_seconds": max_seconds,
                "total_processed": int(row.get("total_processed") or 0),
                "sla_status": status,
            }
        )

    return {
        "items": items,
        "meta": {
            "period_days": days,
        },
    }


@router.get("/api/admin/citizens")
async def admin_list_citizens(
    q: str | None = None,
    name: str | None = None,
    bi: str | None = None,
    email: str | None = None,
    phone: str | None = None,
    status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    limit: int = 20,
    offset: int = 0,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    q_like = f"%{q}%" if q else None
    name_like = f"%{name}%" if name else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    phone_like = f"%{phone}%" if phone else None
    is_active = _parse_is_active(status)
    created_from = _parse_date(date_from)
    created_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "name": name,
        "name_like": name_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "phone": phone,
        "phone_like": phone_like,
        "is_active": is_active,
        "created_from": created_from,
        "created_to": created_to,
        "limit": limit,
        "offset": offset,
    }
    count_result = await db.execute(
        text("""
            SELECT count(*)
            FROM citizenship_citizens
            WHERE (:name IS NULL OR name ILIKE :name_like)
              AND (:bi IS NULL OR bi_number ILIKE :bi_like)
              AND (:email IS NULL OR email ILIKE :email_like)
              AND (:phone IS NULL OR phone ILIKE :phone_like)
              AND (:is_active IS NULL OR is_active = :is_active)
              AND (:created_from IS NULL OR created_at >= :created_from)
              AND (:created_to IS NULL OR created_at <= :created_to)
              AND (
                :q IS NULL
                OR name ILIKE :q_like
                OR email ILIKE :q_like
                OR bi_number ILIKE :q_like
                OR phone ILIKE :q_like
              )
        """),
        params,
    )
    total = count_result.scalar_one()
    result = await db.execute(
        text("""
            SELECT id, name, email, phone, address, is_active, created_at, updated_at,
                   bi_number, birth_date, birth_location_id, residence_location_id, user_id
                      FROM citizenship_citizens
            WHERE (:name IS NULL OR name ILIKE :name_like)
              AND (:bi IS NULL OR bi_number ILIKE :bi_like)
              AND (:email IS NULL OR email ILIKE :email_like)
              AND (:phone IS NULL OR phone ILIKE :phone_like)
              AND (:is_active IS NULL OR is_active = :is_active)
              AND (:created_from IS NULL OR created_at >= :created_from)
              AND (:created_to IS NULL OR created_at <= :created_to)
              AND (
                :q IS NULL
                OR name ILIKE :q_like
                OR email ILIKE :q_like
                OR bi_number ILIKE :q_like
                OR phone ILIKE :q_like
              )
            ORDER BY created_at DESC NULLS LAST
            LIMIT :limit OFFSET :offset
        """),
        params,
    )
    items = [_map_citizen_row(dict(row)) for row in result.mappings().all()]
    return {"items": items, "total": total, "limit": limit, "offset": offset}


@router.get("/api/admin/citizens/export")
async def admin_export_citizens(
    q: str | None = None,
    name: str | None = None,
    bi: str | None = None,
    email: str | None = None,
    phone: str | None = None,
    status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    columns: str | None = None,
    format: str | None = None,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> Response:
    _assert_admin(current_user)
    name_like = f"%{name}%" if name else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    phone_like = f"%{phone}%" if phone else None
    q_like = f"%{q}%" if q else None
    is_active = _parse_is_active(status)
    created_from = _parse_date(date_from)
    created_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "name": name,
        "name_like": name_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "phone": phone,
        "phone_like": phone_like,
        "is_active": is_active,
        "created_from": created_from,
        "created_to": created_to,
    }
    result = await db.execute(
        text("""
            SELECT id, name, email, phone, address, is_active, created_at, updated_at,
                   bi_number, birth_date, birth_location_id, residence_location_id, user_id
            FROM citizenship_citizens
            WHERE (:name IS NULL OR name ILIKE :name_like)
              AND (:bi IS NULL OR bi_number ILIKE :bi_like)
              AND (:email IS NULL OR email ILIKE :email_like)
              AND (:phone IS NULL OR phone ILIKE :phone_like)
              AND (:is_active IS NULL OR is_active = :is_active)
              AND (:created_from IS NULL OR created_at >= :created_from)
              AND (:created_to IS NULL OR created_at <= :created_to)
              AND (
                :q IS NULL
                OR name ILIKE :q_like
                OR email ILIKE :q_like
                OR bi_number ILIKE :q_like
                OR phone ILIKE :q_like
              )
            ORDER BY created_at DESC NULLS LAST
        """),
        params,
    )
    rows = [_map_citizen_row(dict(row)) for row in result.mappings().all()]
    import csv
    import io

    output = io.StringIO()
    allowed_columns = [
        "id",
        "full_name",
        "bi_number",
        "birth_date",
        "email",
        "phone",
        "is_active",
        "status",
        "created_at",
        "updated_at",
    ]
    default_columns = [
        "full_name",
        "bi_number",
        "birth_date",
        "email",
        "phone",
        "status",
        "created_at",
    ]
    fieldnames = _parse_columns(columns, allowed_columns, default_columns)
    export_format = (format or "csv").lower()
    if export_format == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Citizens"
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(key) for key in fieldnames])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"citizens_{datetime.utcnow().date().isoformat()}.xlsx"
        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) for key in fieldnames})
    filename = f"citizens_{datetime.utcnow().date().isoformat()}.csv"
    csv_content = "\ufeff" + output.getvalue()
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/api/admin/exports/presets")
async def admin_get_export_presets(
    module: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    presets = await _load_export_presets(db, current_user, module)
    return {"module": module, "presets": presets}


@router.post("/api/admin/exports/presets")
async def admin_save_export_preset(
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    payload = await request.json()
    module = payload.get("module")
    name = payload.get("name")
    columns = payload.get("columns") or []
    if not module or not name or not isinstance(columns, list):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid preset payload"
        )
    presets = await _save_export_preset(db, current_user, module, name, columns)
    return {"module": module, "presets": presets}


@router.delete("/api/admin/exports/presets")
async def admin_delete_export_preset(
    module: str,
    name: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    presets = await _delete_export_preset(db, current_user, module, name)
    return {"module": module, "presets": presets}


@router.post("/api/admin/exports/citizens")
async def admin_export_citizens_job(
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    payload = await request.json()
    q = payload.get("q")
    name = payload.get("name")
    bi = payload.get("bi")
    email = payload.get("email")
    phone = payload.get("phone")
    status_value = payload.get("status")
    date_from = payload.get("date_from")
    date_to = payload.get("date_to")
    columns = payload.get("columns")
    export_format = (payload.get("format") or "csv").lower()
    name_like = f"%{name}%" if name else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    phone_like = f"%{phone}%" if phone else None
    q_like = f"%{q}%" if q else None
    is_active = _parse_is_active(status_value)
    created_from = _parse_date(date_from)
    created_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "name": name,
        "name_like": name_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "phone": phone,
        "phone_like": phone_like,
        "is_active": is_active,
        "created_from": created_from,
        "created_to": created_to,
    }
    user_id = current_user.get("id") or current_user.get("sub") or ""
    job_payload = {
        "params": params,
        "columns": columns,
        "format": export_format,
    }
    job_id = await _create_export_job(db, "citizens", str(user_id), job_payload)
    from core.exports.tasks import run_export_job

    run_export_job.delay(job_id)
    return {
        "job_id": job_id,
        "status": "pending",
        "status_url": f"/api/admin/exports/jobs/{job_id}",
        "download_url": f"/api/admin/exports/jobs/{job_id}/download",
    }


@router.get("/api/admin/citizens/{citizen_id}")
async def admin_get_citizen(
    citizen_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    result = await db.execute(
        text("""
            SELECT id, name, email, phone, address, is_active, created_at, updated_at,
                   bi_number, birth_date, birth_location_id, residence_location_id, user_id
            FROM citizenship_citizens
            WHERE id = :citizen_id
            LIMIT 1
        """),
        {"citizen_id": citizen_id},
    )
    row = result.mappings().first()
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Citizen not found")
    return _map_citizen_row(dict(row))


@router.get("/api/admin/documents")
async def admin_list_documents(
    q: str | None = None,
    document_type: str | None = None,
    bi: str | None = None,
    email: str | None = None,
    status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    limit: int = 20,
    offset: int = 0,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    q_like = f"%{q}%" if q else None
    type_like = f"%{document_type}%" if document_type else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    status_norm = status.strip().lower() if status else None
    issued_from = _parse_date(date_from)
    issued_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "document_type": document_type,
        "document_type_like": type_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "status": status_norm,
        "issued_from": issued_from,
        "issued_to": issued_to,
        "limit": limit,
        "offset": offset,
    }
    count_result = await db.execute(
        text("""
            SELECT count(*)
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE (:document_type IS NULL OR d.document_type ILIKE :document_type_like)
              AND (:bi IS NULL OR c.bi_number ILIKE :bi_like)
              AND (:email IS NULL OR c.email ILIKE :email_like)
              AND (:issued_from IS NULL OR d.issued_at >= :issued_from)
              AND (:issued_to IS NULL OR d.issued_at <= :issued_to)
              AND (
                :status IS NULL
                OR (:status = 'active' AND (d.valid_until IS NULL OR d.valid_until >= NOW()))
                OR (:status = 'expired' AND d.valid_until < NOW())
                OR (:status = 'unknown' AND d.valid_until IS NULL)
              )
              AND (
                :q IS NULL
                OR d.document_type ILIKE :q_like
                OR c.name ILIKE :q_like
                OR c.bi_number ILIKE :q_like
                OR c.email ILIKE :q_like
              )
        """),
        params,
    )
    total = count_result.scalar_one()
    result = await db.execute(
        text("""
            SELECT d.id,
                   d.citizen_id,
                   d.document_type,
                   d.file_url,
                   d.issued_at,
                   d.valid_until,
                   c.name AS citizen_name,
                   c.bi_number AS bi_number,
                   c.email AS citizen_email
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE (:document_type IS NULL OR d.document_type ILIKE :document_type_like)
              AND (:bi IS NULL OR c.bi_number ILIKE :bi_like)
              AND (:email IS NULL OR c.email ILIKE :email_like)
              AND (:issued_from IS NULL OR d.issued_at >= :issued_from)
              AND (:issued_to IS NULL OR d.issued_at <= :issued_to)
              AND (
                :status IS NULL
                OR (:status = 'active' AND (d.valid_until IS NULL OR d.valid_until >= NOW()))
                OR (:status = 'expired' AND d.valid_until < NOW())
                OR (:status = 'unknown' AND d.valid_until IS NULL)
              )
              AND (
                :q IS NULL
                OR d.document_type ILIKE :q_like
                OR c.name ILIKE :q_like
                OR c.bi_number ILIKE :q_like
                OR c.email ILIKE :q_like
              )
            ORDER BY d.issued_at DESC NULLS LAST
            LIMIT :limit OFFSET :offset
        """),
        params,
    )
    items = [_map_document_row(dict(row)) for row in result.mappings().all()]
    return {"items": items, "total": total, "limit": limit, "offset": offset}


@router.get("/api/admin/documents/export")
async def admin_export_documents(
    q: str | None = None,
    document_type: str | None = None,
    bi: str | None = None,
    email: str | None = None,
    status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    columns: str | None = None,
    format: str | None = None,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> Response:
    _assert_admin(current_user)
    q_like = f"%{q}%" if q else None
    type_like = f"%{document_type}%" if document_type else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    status_norm = status.strip().lower() if status else None
    issued_from = _parse_date(date_from)
    issued_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "document_type": document_type,
        "document_type_like": type_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "status": status_norm,
        "issued_from": issued_from,
        "issued_to": issued_to,
    }
    result = await db.execute(
        text("""
            SELECT d.id,
                   d.citizen_id,
                   d.document_type,
                   d.file_url,
                   d.issued_at,
                   d.valid_until,
                   c.name AS citizen_name,
                   c.bi_number AS bi_number,
                   c.email AS citizen_email
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE (:document_type IS NULL OR d.document_type ILIKE :document_type_like)
              AND (:bi IS NULL OR c.bi_number ILIKE :bi_like)
              AND (:email IS NULL OR c.email ILIKE :email_like)
              AND (:issued_from IS NULL OR d.issued_at >= :issued_from)
              AND (:issued_to IS NULL OR d.issued_at <= :issued_to)
              AND (
                :status IS NULL
                OR (:status = 'active' AND (d.valid_until IS NULL OR d.valid_until >= NOW()))
                OR (:status = 'expired' AND d.valid_until < NOW())
                OR (:status = 'unknown' AND d.valid_until IS NULL)
              )
              AND (
                :q IS NULL
                OR d.document_type ILIKE :q_like
                OR c.name ILIKE :q_like
                OR c.bi_number ILIKE :q_like
                OR c.email ILIKE :q_like
              )
            ORDER BY d.issued_at DESC NULLS LAST
        """),
        params,
    )
    rows = [_map_document_row(dict(row)) for row in result.mappings().all()]
    import csv
    import io

    output = io.StringIO()
    allowed_columns = [
        "id",
        "document_type",
        "citizen_name",
        "citizen_bi",
        "citizen_email",
        "issued_at",
        "valid_until",
        "status",
        "file_url",
        "citizen_id",
    ]
    default_columns = [
        "document_type",
        "citizen_name",
        "citizen_bi",
        "issued_at",
        "valid_until",
        "status",
        "file_url",
    ]
    fieldnames = _parse_columns(columns, allowed_columns, default_columns)
    export_format = (format or "csv").lower()
    if export_format == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Documents"
        ws.append(fieldnames)
        for row in rows:
            ws.append([row.get(key) for key in fieldnames])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"documents_{datetime.utcnow().date().isoformat()}.xlsx"
        return Response(
            content=buffer.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) for key in fieldnames})
    filename = f"documents_{datetime.utcnow().date().isoformat()}.csv"
    csv_content = "\ufeff" + output.getvalue()
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/api/admin/exports/documents")
async def admin_export_documents_job(
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    payload = await request.json()
    q = payload.get("q")
    document_type = payload.get("document_type")
    bi = payload.get("bi")
    email = payload.get("email")
    status_value = payload.get("status")
    date_from = payload.get("date_from")
    date_to = payload.get("date_to")
    columns = payload.get("columns")
    export_format = (payload.get("format") or "csv").lower()
    q_like = f"%{q}%" if q else None
    type_like = f"%{document_type}%" if document_type else None
    bi_like = f"%{bi}%" if bi else None
    email_like = f"%{email}%" if email else None
    status_norm = status_value.strip().lower() if status_value else None
    issued_from = _parse_date(date_from)
    issued_to = _parse_date(date_to, end=True)
    params = {
        "q": q,
        "q_like": q_like,
        "document_type": document_type,
        "document_type_like": type_like,
        "bi": bi,
        "bi_like": bi_like,
        "email": email,
        "email_like": email_like,
        "status": status_norm,
        "issued_from": issued_from,
        "issued_to": issued_to,
    }
    user_id = current_user.get("id") or current_user.get("sub") or ""
    job_payload = {
        "params": params,
        "columns": columns,
        "format": export_format,
    }
    job_id = await _create_export_job(db, "documents", str(user_id), job_payload)
    from core.exports.tasks import run_export_job

    run_export_job.delay(job_id)
    return {
        "job_id": job_id,
        "status": "pending",
        "status_url": f"/api/admin/exports/jobs/{job_id}",
        "download_url": f"/api/admin/exports/jobs/{job_id}/download",
    }


@router.post("/api/admin/exports/preview")
async def admin_export_preview(
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    payload = await request.json()
    if not isinstance(payload, dict):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid payload")
    module = (payload.get("module") or "").strip().lower()
    params_raw = payload.get("params") or {}
    if not isinstance(params_raw, dict):
        params_raw = {}
    q = params_raw.get("q")
    q_like = f"%{q}%" if q else None
    if module == "documents":
        document_type = params_raw.get("document_type")
        bi = params_raw.get("bi")
        email = params_raw.get("email")
        status_value = params_raw.get("status")
        issued_from = _parse_date(params_raw.get("date_from"))
        issued_to = _parse_date(params_raw.get("date_to"), end=True)
        params = {
            "q": q,
            "q_like": q_like,
            "document_type": document_type,
            "document_type_like": f"%{document_type}%" if document_type else None,
            "bi": bi,
            "bi_like": f"%{bi}%" if bi else None,
            "email": email,
            "email_like": f"%{email}%" if email else None,
            "status": status_value.strip().lower() if status_value else None,
            "issued_from": issued_from,
            "issued_to": issued_to,
        }
        count = await _count_documents_export(db, params)
        return {"module": module, "count": count}
    if module == "citizens":
        name = params_raw.get("name")
        bi = params_raw.get("bi")
        email = params_raw.get("email")
        phone = params_raw.get("phone")
        status_value = params_raw.get("status")
        created_from = _parse_date(params_raw.get("date_from"))
        created_to = _parse_date(params_raw.get("date_to"), end=True)
        params = {
            "q": q,
            "q_like": q_like,
            "name": name,
            "name_like": f"%{name}%" if name else None,
            "bi": bi,
            "bi_like": f"%{bi}%" if bi else None,
            "email": email,
            "email_like": f"%{email}%" if email else None,
            "phone": phone,
            "phone_like": f"%{phone}%" if phone else None,
            "is_active": _parse_is_active(status_value),
            "created_from": created_from,
            "created_to": created_to,
        }
        count = await _count_citizens_export(db, params)
        return {"module": module, "count": count}
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unknown module")


@router.post("/api/admin/exports/preview/detail")
async def admin_export_preview_detail(
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    payload = await request.json()
    if not isinstance(payload, dict):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid payload")
    module = (payload.get("module") or "").strip().lower()
    params_raw = payload.get("params") or {}
    if not isinstance(params_raw, dict):
        params_raw = {}
    q = params_raw.get("q")
    q_like = f"%{q}%" if q else None
    if module == "documents":
        document_type = params_raw.get("document_type")
        bi = params_raw.get("bi")
        email = params_raw.get("email")
        status_value = params_raw.get("status")
        issued_from = _parse_date(params_raw.get("date_from"))
        issued_to = _parse_date(params_raw.get("date_to"), end=True)
        params = {
            "q": q,
            "q_like": q_like,
            "document_type": document_type,
            "document_type_like": f"%{document_type}%" if document_type else None,
            "bi": bi,
            "bi_like": f"%{bi}%" if bi else None,
            "email": email,
            "email_like": f"%{email}%" if email else None,
            "status": status_value.strip().lower() if status_value else None,
            "issued_from": issued_from,
            "issued_to": issued_to,
        }
        total = await _count_documents_export(db, params)
        by_status = await _count_documents_by_status(db, params)
        return {"module": module, "count": total, "by_status": by_status}
    if module == "citizens":
        name = params_raw.get("name")
        bi = params_raw.get("bi")
        email = params_raw.get("email")
        phone = params_raw.get("phone")
        status_value = params_raw.get("status")
        created_from = _parse_date(params_raw.get("date_from"))
        created_to = _parse_date(params_raw.get("date_to"), end=True)
        params = {
            "q": q,
            "q_like": q_like,
            "name": name,
            "name_like": f"%{name}%" if name else None,
            "bi": bi,
            "bi_like": f"%{bi}%" if bi else None,
            "email": email,
            "email_like": f"%{email}%" if email else None,
            "phone": phone,
            "phone_like": f"%{phone}%" if phone else None,
            "is_active": _parse_is_active(status_value),
            "created_from": created_from,
            "created_to": created_to,
        }
        total = await _count_citizens_export(db, params)
        by_status = await _count_citizens_by_status(db, params)
        return {"module": module, "count": total, "by_status": by_status}
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unknown module")


@router.get("/api/admin/exports/jobs")
async def admin_export_jobs_list(
    module: str | None = None,
    status_filter: str | None = Query(default=None, alias="status"),
    q: str | None = None,
    job_id: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    limit: int = 20,
    offset: int = 0,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    safe_limit = max(1, min(int(limit or 20), 100))
    safe_offset = max(0, int(offset or 0))
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    module_norm = module.strip().lower() if module else None
    status_norm = status_filter.strip().lower() if status_filter else None
    q_norm = q.strip() if q else None
    job_norm = job_id.strip() if job_id else None
    from_ts = _parse_datetime(date_from)
    to_ts = _parse_datetime(date_to)
    items, total, avg_map, status_counts = await _list_export_jobs(
        db,
        user_id,
        module_norm,
        status_norm,
        q_norm,
        job_norm,
        from_ts,
        to_ts,
        safe_limit,
        safe_offset,
    )
    return {
        "items": items,
        "total": total,
        "limit": safe_limit,
        "offset": safe_offset,
        "averages": avg_map,
        "status_counts": status_counts,
    }


@router.get("/api/admin/exports/jobs/{job_id}/detail")
async def admin_export_job_detail(
    job_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    job = await _fetch_export_job(db, job_id)
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export job not found")
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
    if job_owner and job_owner != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    payload = job.get("request_payload")
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except Exception:
            payload = payload
    logs = await db.execute(
        text("""
            SELECT level, message, created_at
            FROM export_job_logs
            WHERE job_id = :job_id
            ORDER BY created_at ASC
        """),
        {"job_id": job_id},
    )
    log_items = []
    for row in logs.mappings().all():
        log_items.append(
            {
                "level": row.get("level"),
                "message": row.get("message"),
                "created_at": _serialize_datetime(row.get("created_at")),
            }
        )
    return {
        "job": _export_job_response(job_id, job),
        "module": job.get("module"),
        "status": job.get("status"),
        "created_at": _serialize_datetime(job.get("created_at")),
        "updated_at": _serialize_datetime(job.get("updated_at")),
        "payload": payload,
        "logs": log_items,
    }


@router.get("/api/admin/exports/jobs/{job_id}/logs/stream")
async def admin_export_job_logs_stream(
    job_id: str,
    request: Request,
    token: str | None = None,
    since: str | None = None,
    authorization: str | None = Header(default=None),
) -> StreamingResponse:
    current_user = await _resolve_stream_user(request, token, authorization)
    _assert_admin(current_user)
    user_id = str(current_user.get("id") or current_user.get("sub") or "")

    async def event_generator():
        last_ts: datetime | None = _parse_datetime(since)
        while True:
            if await request.is_disconnected():
                break
            async with AsyncSessionLocal() as session:
                job = await _fetch_export_job(session, job_id)
                if not job:
                    payload = {"job_id": job_id, "status": "not_found"}
                    yield f"data: {json.dumps(payload)}\n\n"
                    break
                job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
                if job_owner and job_owner != user_id:
                    payload = {"job_id": job_id, "status": "forbidden"}
                    yield f"data: {json.dumps(payload)}\n\n"
                    break
                logs = await session.execute(
                    text("""
                        SELECT level, message, created_at
                        FROM export_job_logs
                        WHERE job_id = :job_id
                          AND (:last_ts IS NULL OR created_at > :last_ts)
                        ORDER BY created_at ASC
                        LIMIT 50
                    """),
                    {"job_id": job_id, "last_ts": last_ts},
                )
                rows = logs.mappings().all()
            if rows:
                for row in rows:
                    payload = {
                        "level": row.get("level"),
                        "message": row.get("message"),
                        "created_at": _serialize_datetime(row.get("created_at")),
                    }
                    yield f"data: {json.dumps(payload)}\n\n"
                last_ts = rows[-1].get("created_at")
            if job.get("status") in {"done", "failed"} and not rows:
                payload = {"status": job.get("status")}
                yield f"data: {json.dumps(payload)}\n\n"
                break
            await asyncio.sleep(1)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={"cache-control": "no-cache", "x-accel-buffering": "no"},
    )


@router.websocket("/api/admin/exports/jobs/{job_id}/logs/ws")
async def admin_export_job_logs_ws(websocket: WebSocket, job_id: str) -> None:
    token = websocket.query_params.get("token")
    since = websocket.query_params.get("since")
    if not token:
        await websocket.close(code=1008)
        return
    try:
        current_user = _decode_token(token)
    except Exception:
        await websocket.close(code=1008)
        return
    try:
        _assert_admin(current_user)
    except HTTPException:
        await websocket.close(code=1008)
        return
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    await websocket.accept()
    last_ts: datetime | None = _parse_datetime(since)
    try:
        while True:
            async with AsyncSessionLocal() as session:
                job = await _fetch_export_job(session, job_id)
                if not job:
                    await websocket.send_json({"status": "not_found"})
                    break
                job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
                if job_owner and job_owner != user_id:
                    await websocket.send_json({"status": "forbidden"})
                    break
                logs = await session.execute(
                    text("""
                        SELECT level, message, created_at
                        FROM export_job_logs
                        WHERE job_id = :job_id
                          AND (:last_ts IS NULL OR created_at > :last_ts)
                        ORDER BY created_at ASC
                        LIMIT 50
                    """),
                    {"job_id": job_id, "last_ts": last_ts},
                )
                rows = logs.mappings().all()
            if rows:
                for row in rows:
                    await websocket.send_json(
                        {
                            "level": row.get("level"),
                            "message": row.get("message"),
                            "created_at": _serialize_datetime(row.get("created_at")),
                        }
                    )
                last_ts = rows[-1].get("created_at")
            if job.get("status") in {"done", "failed"} and not rows:
                await websocket.send_json({"status": job.get("status")})
                break
            await asyncio.sleep(1)
    except WebSocketDisconnect:
        return
    finally:
        await websocket.close()


@router.get("/api/admin/exports/logs/stream")
async def admin_export_logs_stream(
    request: Request,
    token: str | None = None,
    module: str | None = None,
    since: str | None = None,
    authorization: str | None = Header(default=None),
) -> StreamingResponse:
    current_user = await _resolve_stream_user(request, token, authorization)
    _assert_admin(current_user)
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    module_norm = module.strip().lower() if module else None

    async def event_generator():
        last_ts: datetime | None = _parse_datetime(since)
        while True:
            if await request.is_disconnected():
                break
            async with AsyncSessionLocal() as session:
                rows = await session.execute(
                    text("""
                        SELECT l.job_id, l.level, l.message, l.created_at, j.module
                        FROM export_job_logs l
                        JOIN export_jobs j ON j.id = l.job_id
                        WHERE j.owner_id::text = :owner_id
                          AND (CAST(:module AS TEXT) IS NULL OR j.module = CAST(:module AS TEXT))
                          AND (:last_ts IS NULL OR l.created_at > :last_ts)
                        ORDER BY l.created_at ASC
                        LIMIT 50
                    """),
                    {"owner_id": user_id, "module": module_norm, "last_ts": last_ts},
                )
                items = rows.mappings().all()
            if items:
                for row in items:
                    payload = {
                        "job_id": str(row.get("job_id")),
                        "module": row.get("module"),
                        "level": row.get("level"),
                        "message": row.get("message"),
                        "created_at": _serialize_datetime(row.get("created_at")),
                    }
                    yield f"data: {json.dumps(payload)}\n\n"
                last_ts = items[-1].get("created_at")
            await asyncio.sleep(1)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={"cache-control": "no-cache", "x-accel-buffering": "no"},
    )


@router.websocket("/api/admin/exports/logs/ws")
async def admin_export_logs_ws(websocket: WebSocket) -> None:
    token = websocket.query_params.get("token")
    module = websocket.query_params.get("module")
    since = websocket.query_params.get("since")
    if not token:
        await websocket.close(code=1008)
        return
    try:
        current_user = _decode_token(token)
    except Exception:
        await websocket.close(code=1008)
        return
    try:
        _assert_admin(current_user)
    except HTTPException:
        await websocket.close(code=1008)
        return
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    module_norm = module.strip().lower() if module else None
    await websocket.accept()
    last_ts: datetime | None = _parse_datetime(since)
    try:
        while True:
            async with AsyncSessionLocal() as session:
                rows = await session.execute(
                    text("""
                        SELECT l.job_id, l.level, l.message, l.created_at, j.module
                        FROM export_job_logs l
                        JOIN export_jobs j ON j.id = l.job_id
                        WHERE j.owner_id::text = :owner_id
                          AND (CAST(:module AS TEXT) IS NULL OR j.module = CAST(:module AS TEXT))
                          AND (:last_ts IS NULL OR l.created_at > :last_ts)
                        ORDER BY l.created_at ASC
                        LIMIT 50
                    """),
                    {"owner_id": user_id, "module": module_norm, "last_ts": last_ts},
                )
                items = rows.mappings().all()
            if items:
                for row in items:
                    await websocket.send_json(
                        {
                            "job_id": str(row.get("job_id")),
                            "module": row.get("module"),
                            "level": row.get("level"),
                            "message": row.get("message"),
                            "created_at": _serialize_datetime(row.get("created_at")),
                        }
                    )
                last_ts = items[-1].get("created_at")
            await asyncio.sleep(1)
    except WebSocketDisconnect:
        return
    finally:
        await websocket.close()


@router.get("/api/admin/exports/jobs/timeline")
async def admin_export_jobs_timeline(
    module: str | None = None,
    days: int = 14,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    module_norm = module.strip().lower() if module else None
    safe_days = max(1, min(int(days or 14), 60))
    rows = await db.execute(
        text("""
            SELECT date_trunc('day', created_at) AS day,
                   AVG(EXTRACT(EPOCH FROM (updated_at - created_at))) AS avg_seconds,
                   COUNT(*) AS total
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND status = 'done'
              AND created_at >= now() - (:days || ' days')::interval
              AND (CAST(:module AS TEXT) IS NULL OR module = CAST(:module AS TEXT))
            GROUP BY date_trunc('day', created_at)
            ORDER BY day ASC
        """),
        {"owner_id": user_id, "module": module_norm, "days": str(safe_days)},
    )
    items = []
    for row in rows.mappings().all():
        items.append(
            {
                "day": _serialize_datetime(row.get("day")),
                "avg_seconds": float(row.get("avg_seconds") or 0.0),
                "total": int(row.get("total") or 0),
            }
        )
    return {"items": items}


@router.get("/api/admin/exports/jobs/timeline/compare")
async def admin_export_jobs_timeline_compare(
    days: int = 14,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    safe_days = max(1, min(int(days or 14), 60))
    rows = await db.execute(
        text("""
            SELECT date_trunc('day', created_at) AS day,
                   module,
                   AVG(EXTRACT(EPOCH FROM (updated_at - created_at))) AS avg_seconds,
                   COUNT(*) AS total
            FROM export_jobs
            WHERE owner_id::text = :owner_id
              AND status = 'done'
              AND created_at >= now() - (:days || ' days')::interval
              AND module IN ('citizens', 'documents')
            GROUP BY date_trunc('day', created_at), module
            ORDER BY day ASC
        """),
        {"owner_id": user_id, "days": str(safe_days)},
    )
    series: dict[str, list[dict[str, Any]]] = {"citizens": [], "documents": []}
    for row in rows.mappings().all():
        module_name = row.get("module")
        if module_name not in series:
            continue
        series[module_name].append(
            {
                "day": _serialize_datetime(row.get("day")),
                "avg_seconds": float(row.get("avg_seconds") or 0.0),
                "total": int(row.get("total") or 0),
            }
        )
    return {"series": series}


@router.post("/api/admin/exports/jobs/{job_id}/repeat")
async def admin_export_job_repeat(
    job_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    job = await _fetch_export_job(db, job_id)
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export job not found")
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
    if job_owner and job_owner != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    payload = job.get("request_payload") or {}
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except Exception:
            payload = {}
    new_job_id = await _create_export_job(db, str(job.get("module") or "exports"), user_id, payload)
    from core.exports.tasks import run_export_job

    run_export_job.delay(new_job_id)
    return {
        "job_id": new_job_id,
        "status": "pending",
        "status_url": f"/api/admin/exports/jobs/{new_job_id}",
        "download_url": f"/api/admin/exports/jobs/{new_job_id}/download",
    }


@router.post("/api/admin/exports/jobs/{job_id}/reprocess")
async def admin_export_job_reprocess(
    job_id: str,
    request: Request,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    job = await _fetch_export_job(db, job_id)
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export job not found")
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
    if job_owner and job_owner != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    base_payload = job.get("request_payload") or {}
    if isinstance(base_payload, str):
        try:
            base_payload = json.loads(base_payload)
        except Exception:
            base_payload = {}
    override_payload = await request.json()
    if not isinstance(override_payload, dict):
        override_payload = {}
    merged_payload = {**base_payload, **override_payload}
    new_job_id = await _create_export_job(
        db, str(job.get("module") or "exports"), user_id, merged_payload
    )
    from core.exports.tasks import run_export_job

    run_export_job.delay(new_job_id)
    return {
        "job_id": new_job_id,
        "status": "pending",
        "status_url": f"/api/admin/exports/jobs/{new_job_id}",
        "download_url": f"/api/admin/exports/jobs/{new_job_id}/download",
    }


@router.get("/api/admin/exports/jobs/{job_id}")
async def admin_export_job_status(
    job_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    job = await _fetch_export_job(db, job_id)
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export job not found")
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
    if job_owner and job_owner != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    return _export_job_response(job_id, job)


@router.get("/api/admin/exports/jobs/{job_id}/stream")
async def admin_export_job_stream(
    job_id: str,
    request: Request,
    token: str | None = None,
    authorization: str | None = Header(default=None),
) -> StreamingResponse:
    current_user = await _resolve_stream_user(request, token, authorization)
    _assert_admin(current_user)
    user_id = str(current_user.get("id") or current_user.get("sub") or "")

    async def event_generator():
        last_status = None
        while True:
            if await request.is_disconnected():
                break
            async with AsyncSessionLocal() as session:
                job = await _fetch_export_job(session, job_id)
            if not job:
                payload = {"job_id": job_id, "status": "not_found"}
                yield f"data: {json.dumps(payload)}\n\n"
                break
            job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
            if job_owner and job_owner != user_id:
                payload = {"job_id": job_id, "status": "forbidden"}
                yield f"data: {json.dumps(payload)}\n\n"
                break
            status_value = job.get("status")
            if status_value != last_status:
                last_status = status_value
                payload = _export_job_response(job_id, job)
                yield f"data: {json.dumps(payload)}\n\n"
                if status_value in {"done", "failed"}:
                    break
            await asyncio.sleep(1)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={"cache-control": "no-cache", "x-accel-buffering": "no"},
    )


@router.websocket("/api/admin/exports/jobs/{job_id}/ws")
async def admin_export_job_ws(websocket: WebSocket, job_id: str) -> None:
    token = websocket.query_params.get("token")
    if not token:
        await websocket.close(code=1008)
        return
    try:
        current_user = _decode_token(token)
    except Exception:
        await websocket.close(code=1008)
        return
    try:
        _assert_admin(current_user)
    except HTTPException:
        await websocket.close(code=1008)
        return
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    await websocket.accept()
    last_status = None
    try:
        while True:
            async with AsyncSessionLocal() as session:
                job = await _fetch_export_job(session, job_id)
            if not job:
                await websocket.send_json({"job_id": job_id, "status": "not_found"})
                break
            job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
            if job_owner and job_owner != user_id:
                await websocket.send_json({"job_id": job_id, "status": "forbidden"})
                break
            status_value = job.get("status")
            if status_value != last_status:
                last_status = status_value
                await websocket.send_json(_export_job_response(job_id, job))
                if status_value in {"done", "failed"}:
                    break
            await asyncio.sleep(1)
    except WebSocketDisconnect:
        return
    finally:
        await websocket.close()


@router.get("/api/admin/exports/jobs/{job_id}/download")
async def admin_export_job_download(
    job_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> Response:
    _assert_admin(current_user)
    job = await _fetch_export_job(db, job_id)
    if not job:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Export job not found")
    user_id = str(current_user.get("id") or current_user.get("sub") or "")
    job_owner = str(job.get("owner_id")) if job.get("owner_id") else None
    if job_owner and job_owner != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    if job.get("status") != "done":
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Export not ready")
    filename = job.get("result_filename", "export")
    data = job.get("result_data") or b""
    if isinstance(data, memoryview):
        data = data.tobytes()
    return Response(
        content=data,
        media_type=job.get("result_content_type") or "application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/api/admin/documents/{document_id}")
async def admin_get_document(
    document_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    _assert_admin(current_user)
    result = await db.execute(
        text("""
            SELECT d.id,
                   d.citizen_id,
                   d.document_type,
                   d.file_url,
                   d.issued_at,
                   d.valid_until,
                   c.name AS citizen_name,
                   c.bi_number AS bi_number,
                   c.email AS citizen_email
            FROM wallet_documents d
            LEFT JOIN citizenship_citizens c ON c.id::text = d.citizen_id
            WHERE d.id = :document_id
            LIMIT 1
        """),
        {"document_id": document_id},
    )
    row = result.mappings().first()
    if not row:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return _map_document_row(dict(row))


@router.get("/api/admin/territory/provinces")
async def admin_list_provinces(
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    _assert_admin(current_user)
    result = await db.execute(
        text("""
            SELECT id, name, code, type, parent_id
            FROM locations
            WHERE type = 'province'
              AND parent_id IS NULL
            ORDER BY name
        """),
    )
    return [_map_territory_row(dict(row)) for row in result.mappings().all()]


@router.get("/api/admin/territory/provinces/{province_id}/municipalities")
async def admin_list_municipalities(
    province_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    _assert_admin(current_user)
    result = await db.execute(
        text("""
            SELECT id, name, code, type, parent_id
            FROM locations
            WHERE type = 'municipality'
              AND parent_id = :province_id::uuid
            ORDER BY name
        """),
        {"province_id": province_id},
    )
    return [_map_territory_row(dict(row)) for row in result.mappings().all()]


@router.get("/api/admin/territory/municipalities/{municipality_id}/communes")
async def admin_list_communes(
    municipality_id: str,
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    _assert_admin(current_user)
    result = await db.execute(
        text("""
            SELECT id, name, code, type, parent_id
            FROM locations
            WHERE type = 'commune'
              AND parent_id = :municipality_id::uuid
            ORDER BY name
        """),
        {"municipality_id": municipality_id},
    )
    return [_map_territory_row(dict(row)) for row in result.mappings().all()]


@router.get("/api/admin/territory/tree")
async def admin_territory_tree(
    current_user: dict[str, Any] = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    _assert_admin(current_user)
    provinces_result = await db.execute(
        text("""
            SELECT id, name, code, type, parent_id
            FROM locations
            WHERE type = 'province'
              AND parent_id IS NULL
            ORDER BY name
        """),
    )
    provinces = [_map_territory_row(dict(row)) for row in provinces_result.mappings().all()]
    if not provinces:
        return []

    province_ids = [row["id"] for row in provinces]
    municipalities_result = await db.execute(
        text("""
            SELECT id, name, code, type, parent_id
            FROM locations
            WHERE type = 'municipality'
              AND parent_id::text IN :province_ids
            ORDER BY name
        """).bindparams(bindparam("province_ids", expanding=True)),
        {"province_ids": province_ids},
    )
    municipalities = [
        _map_territory_row(dict(row)) for row in municipalities_result.mappings().all()
    ]
    municipality_ids = [row["id"] for row in municipalities]
    communes: list[dict[str, Any]] = []
    if municipality_ids:
        communes_result = await db.execute(
            text("""
                SELECT id, name, code, type, parent_id
                FROM locations
                WHERE type = 'commune'
                  AND parent_id::text IN :municipality_ids
                ORDER BY name
            """).bindparams(bindparam("municipality_ids", expanding=True)),
            {"municipality_ids": municipality_ids},
        )
        communes = [_map_territory_row(dict(row)) for row in communes_result.mappings().all()]

    municipalities_by_parent: dict[str, list[dict[str, Any]]] = {}
    for item in municipalities:
        parent_id = item.get("parent_id")
        if not parent_id:
            continue
        municipalities_by_parent.setdefault(parent_id, []).append(item)

    communes_by_parent: dict[str, list[dict[str, Any]]] = {}
    for item in communes:
        parent_id = item.get("parent_id")
        if not parent_id:
            continue
        communes_by_parent.setdefault(parent_id, []).append(item)

    tree: list[dict[str, Any]] = []
    for province in provinces:
        province_children = []
        for municipality in municipalities_by_parent.get(province["id"], []):
            municipality_children = communes_by_parent.get(municipality["id"], [])
            province_children.append({**municipality, "children": municipality_children})
        tree.append({**province, "children": province_children})
    return tree


@router.post("/auth/refresh")
@router.post("/api/auth/refresh")
@router.post("/api/v1/auth/refresh")
async def auth_refresh(request: Request) -> dict[str, Any]:
    from jwt.exceptions import InvalidTokenError

    payload = await _read_login_payload(request)
    refresh_token = (payload.get("refresh_token") or "").strip()
    if not refresh_token:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Missing refresh token")
    try:
        jwt_handler = JWTHandler(secret_key=settings.SECRET_KEY)
        claims = jwt_handler.decode_token(refresh_token)
    except InvalidTokenError:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid refresh token"
        ) from None
    if claims.get("type") != "refresh":
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid refresh token"
        )
    access_token = _issue_access_token(claims)
    new_refresh_token = _issue_refresh_token(claims)
    return {
        "access_token": access_token,
        "refresh_token": new_refresh_token,
        "token_type": "bearer",
    }


@router.post("/auth/logout")
@router.post("/api/auth/logout")
@router.post("/api/v1/auth/logout")
async def auth_logout() -> dict[str, str]:
    return {"status": "ok"}


@router.get("/api/debug/auth-check")
async def debug_auth_check(
    request: Request,
    db: AsyncSession = Depends(get_db),
    authorization: str | None = Header(default=None),
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "timestamp": time.time(),
        "client_ip": request.client.host if request.client else None,
        "db_ok": False,
    }
    try:
        await db.execute(text("SELECT 1"))
        result["db_ok"] = True
    except Exception as exc:
        await db.rollback()
        result["db_error"] = str(exc)
    try:
        result["iam_users_count"] = (
            await db.execute(text("SELECT count(*) FROM iam_users"))
        ).scalar_one()
    except Exception as exc:
        await db.rollback()
        result["iam_users_error"] = str(exc)
    try:
        rows = await db.execute(
            text(
                "\n                SELECT email, username, status, is_active\n                FROM iam_users\n                ORDER BY created_at DESC NULLS LAST\n                LIMIT 10\n                "
            )
        )
        result["iam_users_sample"] = [dict(row) for row in rows.mappings().all()]
    except Exception:
        await db.rollback()
        result["iam_users_sample"] = []
    if authorization:
        try:
            user = await get_current_user(authorization)
            result["user"] = user
            role = user.get("role") or _map_primary_role(user.get("roles", []), user.get("email"))
            if role == "CITIZEN":
                result["recommended_dashboard"] = "/citizen/portal"
            else:
                result["recommended_dashboard"] = "/admin"
        except HTTPException as exc:
            result["auth_error"] = exc.detail
    return result


@router.get("/api/v1/identidade/bi/tipos-evento")
def bi_event_types() -> list[dict[str, str]]:
    return [
        {"code": "EMISSAO", "label": "Emissao"},
        {"code": "RENOVACAO", "label": "Renovacao"},
        {"code": "SEGUNDA_VIA", "label": "Segunda Via"},
    ]


@router.post("/api/v1/identidade/bi/emit")
def bi_emit(
    payload: BIEmitPayload,
    authorization: str | None = Header(default=None),
    x_token: str | None = Header(default=None, alias="X-Token"),
) -> dict[str, Any]:
    if not authorization and (not x_token):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="missing token")
    return {
        "status": "queued",
        "request_id": str(uuid.uuid4()),
        "citizen_fuc_id": payload.citizen_fuc_id,
    }


@router.get("/api/v1/services")
def list_services() -> list[dict[str, Any]]:
    return _SERVICES


def _catalog_from_sla() -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for code, data in SLA_DEFINITIONS.items():
        if code == "DEFAULT":
            continue
        items.append(
            {
                "code": code,
                "name": data.get("description", code),
                "category": code.split("_")[0].lower(),
                "description": data.get("description", ""),
                "price": 0.0,
            }
        )
    return items


async def _catalog_from_service_catalog(
    db: AsyncSession, limit: int, offset: int, query: str | None
) -> list[dict[str, Any]]:
    where = ""
    params: dict[str, Any] = {"limit": limit, "offset": offset}
    if query:
        where = "WHERE service_code ILIKE :q OR name ILIKE :q OR category ILIKE :q"
        params["q"] = f"%{query}%"
    try:
        rows = await db.execute(
            text(f"""
            SELECT service_code, name, category, description, price
            FROM service_catalog
            WHERE is_active = true
            {("AND " + where[6:]) if where else ""}
            ORDER BY name
            LIMIT :limit OFFSET :offset
        """),
            params,
        )
        return [
            {
                "code": row.service_code,
                "name": row.name,
                "category": row.category,
                "description": row.description,
                "price": float(row.price or 0.0),
            }
            for row in rows.mappings().all()
        ]
    except Exception:
        await db.rollback()
        return []


async def _catalog_from_sla_base(
    db: AsyncSession, limit: int, offset: int, query: str | None
) -> list[dict[str, Any]]:
    where = ""
    params: dict[str, Any] = {"limit": limit, "offset": offset}
    if query:
        where = "WHERE service_id ILIKE :q OR service_name ILIKE :q OR module ILIKE :q"
        params["q"] = f"%{query}%"
    try:
        rows = await db.execute(
            text(f"""
            SELECT service_id, service_name, module
            FROM sla_base
            {where}
            ORDER BY service_name
            LIMIT :limit OFFSET :offset
        """),
            params,
        )
        return [
            {
                "code": row.service_id,
                "name": row.service_name,
                "category": row.module,
                "description": row.service_name,
                "price": 0.0,
            }
            for row in rows.mappings().all()
        ]
    except Exception:
        await db.rollback()
        return []


def _resolve_catalog_items() -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    if _SERVICE_CATALOG:
        items.extend(_SERVICE_CATALOG)
    items.extend(_catalog_from_sla())
    if _SERVICES:
        items.extend(
            {
                "code": service.get("code"),
                "name": service.get("name"),
                "category": service.get("id"),
                "description": service.get("name"),
                "price": service.get("price", 0.0),
            }
            for service in _SERVICES
        )
    deduped: dict[str, dict[str, Any]] = {}
    for item in items:
        code = str(item.get("code") or "").upper()
        if not code:
            continue
        if code not in deduped:
            deduped[code] = item
    return list(deduped.values())


def _normalize_module_name(module_code: str) -> str:
    if not module_code:
        return "Outros"
    return _MODULE_LABELS.get(module_code, module_code.replace("_", " ").title())


def _load_module_yaml() -> dict[str, dict[str, Any]]:
    modules: dict[str, dict[str, Any]] = {}
    if not _MODULES_PATH.exists():
        return modules
    for module_dir in _MODULES_PATH.iterdir():
        if not module_dir.is_dir():
            continue
        module_yaml = module_dir / "module.yaml"
        if not module_yaml.exists():
            continue
        try:
            with module_yaml.open("r", encoding="utf-8") as handle:
                data = yaml.safe_load(handle) or {}
            code = data.get("code") or data.get("module") or module_dir.name
            modules[code] = data
        except Exception:
            continue
    return modules


def _module_key_from_item(item: dict[str, Any]) -> str:
    raw = (item.get("category") or "").strip().lower()
    if not raw:
        raw = item.get("code", "").split("_")[0].lower() if item.get("code") else "outros"
    return _CATEGORY_MODULE_MAP.get(raw, raw)


def _sla_payload(service_code: str) -> dict[str, Any]:
    sla = get_sla_for_service(service_code)
    seconds = float(sla.get("target_sla_seconds") or 0)
    hours = round(seconds / 3600, 2) if seconds else None
    days = round(seconds / 86400, 2) if seconds else None
    return {
        "sla_seconds": seconds,
        "sla_hours": hours,
        "sla_days": days,
    }


def _build_portal_catalog(
    items: list[dict[str, Any]],
    module_yaml: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    modules: dict[str, dict[str, Any]] = {}
    for item in items:
        module_key = _module_key_from_item(item)
        module_data = modules.get(module_key)
        if not module_data:
            yaml_data = module_yaml.get(module_key, {})
            ui_override = _MODULE_UI_OVERRIDES.get(module_key, {})
            ui = {**ui_override, **(yaml_data.get("ui") or {})}
            module_data = {
                "module": module_key,
                "name": _normalize_module_name(module_key),
                "description": yaml_data.get("description"),
                "ui": ui,
                "services": [],
            }
            modules[module_key] = module_data
        service_code = item.get("code")
        sla_info = _sla_payload(service_code) if service_code else {}
        module_data["services"].append(
            {
                "id": service_code,
                "code": service_code,
                "name": item.get("name"),
                "description": item.get("description"),
                "price": item.get("price"),
                "category": item.get("category"),
                **sla_info,
            }
        )

    # Merge services defined only in module.yaml (if any)
    for module_key, yaml_data in module_yaml.items():
        module_data = modules.get(module_key)
        if module_data is None:
            services = yaml_data.get("services") or []
            if not services:
                continue
            ui_override = _MODULE_UI_OVERRIDES.get(module_key, {})
            ui = {**ui_override, **(yaml_data.get("ui") or {})}
            modules[module_key] = {
                "module": module_key,
                "name": _normalize_module_name(module_key),
                "description": yaml_data.get("description"),
                "ui": ui,
                "services": [
                    {
                        "id": service.get("id") or service.get("code"),
                        "code": service.get("id") or service.get("code"),
                        "name": service.get("name"),
                        "description": service.get("description"),
                        "price": service.get("price"),
                        **_sla_payload(service.get("id") or service.get("code") or ""),
                    }
                    for service in services
                ],
            }
            continue
        services = yaml_data.get("services") or []
        if not services:
            continue
        existing_codes = {
            (service.get("code") or service.get("id") or "").upper()
            for service in module_data["services"]
        }
        for service in services:
            service_code = service.get("id") or service.get("code")
            if not service_code:
                continue
            if service_code.upper() in existing_codes:
                continue
            module_data["services"].append(
                {
                    "id": service_code,
                    "code": service_code,
                    "name": service.get("name"),
                    "description": service.get("description"),
                    "price": service.get("price"),
                    "category": service.get("category") or module_key,
                    **_sla_payload(service_code),
                }
            )

    for module in modules.values():
        module["services"] = sorted(
            module["services"],
            key=lambda service: (service.get("name") or "").lower(),
        )

    def _order_key(module: dict[str, Any]) -> tuple[int, str]:
        order_value = module.get("ui", {}).get("order")
        if isinstance(order_value, (int, float)):
            order = int(order_value)
        else:
            order = 999
        return (order, module.get("name") or "")

    ordered = sorted(modules.values(), key=_order_key)
    return ordered


def _default_form_for(service_code: str, service_name: str) -> list[CatalogField]:
    lower = service_code.lower()
    fields: list[CatalogField] = [
        CatalogField(
            key="citizen_name",
            label="Nome completo",
            required=True,
            placeholder="Nome conforme documento",
        ),
        CatalogField(
            key="citizen_document",
            label="Documento de identificacao",
            required=True,
            placeholder="BI ou Passaporte",
        ),
    ]
    if "nif" in lower or "tax" in lower:
        fields.append(
            CatalogField(
                key="nif",
                label="Numero de NIF",
                type="nif",
                mask="nif",
                placeholder="Opcional se for emissao",
            )
        )
    if "bi" in lower or "identity" in lower:
        fields.append(
            CatalogField(key="bi_number", label="Numero BI", type="bi", mask="bi", required=True)
        )
    if "certidao" in lower or "registry" in lower or "civil" in lower:
        fields.append(
            CatalogField(
                key="certificate_type",
                label="Tipo de certidao",
                type="select",
                required=True,
                options=["nascimento", "casamento", "obito", "outra"],
            )
        )
    if "agua" in lower or "water" in lower or "energia" in lower or "energy" in lower:
        fields.append(
            CatalogField(
                key="address", label="Morada", required=True, placeholder="Rua, bairro, municipio"
            )
        )
    fields.append(
        CatalogField(
            key="delivery_mode",
            label="Entrega",
            type="radio",
            required=True,
            options=["digital", "presencial", "expresso"],
        )
    )
    fields.append(
        CatalogField(key="supporting_files", label="Anexos", type="upload", required=False)
    )
    fields.append(
        CatalogField(
            key="terms_accept",
            label="Confirmo que os dados sao verdadeiros",
            type="checkbox",
            required=True,
        )
    )
    fields.append(
        CatalogField(
            key="notes",
            label="Observacoes",
            type="textarea",
            placeholder=f"Detalhes adicionais para {service_name}",
        )
    )
    return fields


def _build_schema_from_fields(fields: list[CatalogField]) -> dict[str, Any]:
    return {
        "version": "1.0",
        "groups": [
            {
                "key": "default",
                "label": "Formulario",
                "fields": [field.model_dump() for field in fields],
            }
        ],
        "validations": [],
        "dependencies": [],
    }


def _find_catalog_item(service_code: str, items: list[dict[str, Any]]) -> dict[str, Any] | None:
    normalized = service_code.lower()
    for item in items:
        code = str(item.get("code", "")).lower()
        if code == normalized:
            return item
    return None


@router.get("/api/v1/service-catalog")
async def list_service_catalog(
    query: str | None = None,
    limit: int = 200,
    offset: int = 0,
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    db_items = await _catalog_from_service_catalog(db, limit, offset, query)
    if db_items:
        return [{**item, **_sla_payload(item.get("code") or "")} for item in db_items]
    db_items = await _catalog_from_sla_base(db, limit, offset, query)
    if db_items:
        return [{**item, **_sla_payload(item.get("code") or "")} for item in db_items]
    items = _resolve_catalog_items()
    if query:
        q = query.lower()
        items = [
            item
            for item in items
            if q in str(item.get("name", "")).lower() or q in str(item.get("code", "")).lower()
        ]
    items = [{**item, **_sla_payload(item.get("code") or "")} for item in items]
    return items[offset : offset + limit]


@router.get("/portal/catalog")
@router.get("/api/v1/portal/catalog")
async def get_portal_catalog(
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    items = await _catalog_from_service_catalog(db, limit=2000, offset=0, query=None)
    if not items:
        items = await _catalog_from_sla_base(db, limit=2000, offset=0, query=None)
    if not items:
        items = _resolve_catalog_items()
    module_yaml = _load_module_yaml()
    return _build_portal_catalog(items, module_yaml)


@router.get("/api/v1/service-catalog/{service_code}")
async def get_service_catalog_item(
    service_code: str,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    db_items = await _catalog_from_service_catalog(db, limit=1, offset=0, query=service_code)
    if db_items:
        return db_items[0]
    db_items = await _catalog_from_sla_base(db, limit=1, offset=0, query=service_code)
    if db_items:
        return db_items[0]
    items = _resolve_catalog_items()
    item = _find_catalog_item(service_code, items)
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="service not found")
    return item


@router.get("/api/v1/service-catalog/{service_code}/form", response_model=CatalogForm)
async def get_service_catalog_form(
    service_code: str,
    db: AsyncSession = Depends(get_db),
) -> CatalogForm:
    item = await get_service_catalog_item(service_code, db)
    try:
        row = await db.execute(
            text("""
                SELECT schema FROM service_forms
                WHERE service_code = :code AND is_active = true
                ORDER BY updated_at DESC
                LIMIT 1
            """),
            {"code": service_code},
        )
        result = row.fetchone()
        if result and result[0]:
            schema = result[0]
            if isinstance(schema, str):
                schema = json.loads(schema)
            if isinstance(schema, list):
                fields = [CatalogField(**field) for field in schema]
                return CatalogForm(code=service_code, form_schema=_build_schema_from_fields(fields))
            if isinstance(schema, dict):
                return CatalogForm(code=service_code, form_schema=schema)
    except Exception:
        await db.rollback()
    fields = _default_form_for(service_code, item.get("name", service_code))
    return CatalogForm(code=service_code, form_schema=_build_schema_from_fields(fields))


@router.post("/api/v1/service-catalog/{service_code}/submit")
async def submit_service_catalog(
    service_code: str,
    payload: dict[str, Any],
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    item = await get_service_catalog_item(service_code, db)
    order_id = str(uuid.uuid4())
    order = {"id": order_id, "service_id": service_code, "status": "created", "payload": payload}
    _ORDERS[order_id] = order
    price = float(item.get("price", 0.0) or 0.0)
    return {"order_id": order_id, "payment_required": price > 0, "price": price}


@router.post(
    "/api/v1/financas/invoices", response_model=InvoiceResponse, status_code=status.HTTP_201_CREATED
)
async def create_financas_invoice(
    payload: CreateInvoiceSchema,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> InvoiceResponse:
    invoice_service, _, _, _ = _get_financas_services(db)
    try:
        invoice = await invoice_service.create_invoice(payload)
        await db.commit()
        return _serialize_invoice(invoice)
    except DomainValidationError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))
    except FUCError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc))


@router.get("/api/v1/financas/invoices", response_model=list[InvoiceResponse])
async def list_financas_invoices(
    citizen_id: str | None = None,
    limit: int = 100,
    offset: int = 0,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[InvoiceResponse]:
    _, _, invoice_repo, _ = _get_financas_services(db)
    if citizen_id:
        invoices = await invoice_repo.get_by_citizen(citizen_id)
    else:
        invoices = await invoice_repo.list_all(limit=limit, offset=offset)
    return [_serialize_invoice(inv) for inv in invoices]


@router.get("/api/v1/financas/invoices/{invoice_id}", response_model=InvoiceResponse)
async def get_financas_invoice(
    invoice_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> InvoiceResponse:
    invoice_service, _, _, _ = _get_financas_services(db)
    try:
        invoice = await invoice_service.get_invoice(invoice_id)
        return _serialize_invoice(invoice)
    except InvoiceNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))


@router.get("/api/v1/financas/invoices/citizen/{citizen_id}", response_model=list[InvoiceResponse])
async def list_financas_invoices_citizen(
    citizen_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[InvoiceResponse]:
    invoice_service, _, _, _ = _get_financas_services(db)
    invoices = await invoice_service.list_citizen_invoices(citizen_id)
    return [_serialize_invoice(inv) for inv in invoices]


@router.post(
    "/api/v1/financas/payments", response_model=PaymentResponse, status_code=status.HTTP_201_CREATED
)
async def register_financas_payment(
    payload: CreatePaymentSchema,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> PaymentResponse:
    _, payment_service, _, _ = _get_financas_services(db)
    try:
        payment = await payment_service.register_payment(payload)
        await db.commit()
        return _serialize_payment(payment)
    except DuplicatePaymentError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(exc))
    except InvoiceNotFoundError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))
    except InvalidInvoiceStateError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))
    except DomainValidationError as exc:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc))


@router.get("/api/v1/financas/payments", response_model=list[PaymentResponse])
async def list_financas_payments(
    limit: int = 100,
    offset: int = 0,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[PaymentResponse]:
    _, _, _, payment_repo = _get_financas_services(db)
    payments = await payment_repo.list_all(limit=limit, offset=offset)
    return [_serialize_payment(pay) for pay in payments]


@router.get("/api/v1/financas/payments/citizen/{citizen_id}", response_model=list[PaymentResponse])
async def list_financas_payments_citizen(
    citizen_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: dict[str, Any] = Depends(get_current_user),
) -> list[PaymentResponse]:
    _, payment_service, _, _ = _get_financas_services(db)
    payments = await payment_service.get_payment_history(citizen_id)
    return [_serialize_payment(pay) for pay in payments]


@router.post("/api/v1/financas/payments/confirm")
async def confirm_financas_payment(payload: dict[str, Any]) -> dict[str, str]:
    return {"status": "received", "message": "confirmation queued"}


@router.post("/api/v1/orders", status_code=status.HTTP_201_CREATED)
def create_order(payload: OrderCreatePayload) -> dict[str, Any]:
    order_id = str(uuid.uuid4())
    order = {"id": order_id, "service_id": payload.service_id, "status": "created"}
    _ORDERS[order_id] = order
    return order


@router.post("/api/v1/orders/{order_id}/documents")
def attach_documents(order_id: str, payload: DocumentsPayload) -> dict[str, Any]:
    order = _ORDERS.get(order_id)
    if not order:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="order not found")
    order["documents"] = [doc.model_dump() for doc in payload.documents]
    return {"status": "ok", "order_id": order_id}


@router.post("/api/v1/orders/{order_id}/submit")
def submit_order(order_id: str) -> dict[str, Any]:
    order = _ORDERS.get(order_id)
    if not order:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="order not found")
    order["status"] = "submitted"
    return {"status": "submitted", "order_id": order_id}


@router.post("/api/v1/payments/{order_id}/generate")
def generate_payment(order_id: str) -> dict[str, Any]:
    order = _ORDERS.get(order_id)
    if not order:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="order not found")
    reference = f"PAY-{uuid.uuid4().hex[:10].upper()}"
    payment = {"reference": reference, "order_id": order_id, "status": "PENDING"}
    _PAYMENTS[reference] = payment
    return payment


@router.post("/api/v1/payments/{reference}/confirm")
def confirm_payment(reference: str) -> dict[str, Any]:
    payment = _PAYMENTS.get(reference)
    if not payment:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="invalid reference")
    payment["status"] = "CONFIRMED"
    return {"reference": reference, "status": payment["status"]}


@router.post("/api/v1/orders/{order_id}/complete")
def complete_order(order_id: str) -> dict[str, Any]:
    order = _ORDERS.get(order_id)
    if not order:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="order not found")
    order["status"] = "completed"
    return {"status": "completed", "order_id": order_id}


@router.get("/api/v1/orders/{order_id}/receipt")
def get_receipt(order_id: str) -> dict[str, Any]:
    order = _ORDERS.get(order_id)
    if not order:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="order not found")
    return {"order_id": order_id, "status": "issued", "items": order.get("documents", [])}


@router.get("/api/v1/{module}/ping")
def module_ping(module: str) -> dict[str, str]:
    if module not in _PING_MODULES:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="module not found")
    return {
        "status": "success",
        "message": "pong",
        "module": module,
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
    }
